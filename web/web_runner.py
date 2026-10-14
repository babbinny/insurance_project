# -*- coding: utf-8 -*-
"""
웹페이지(web/index.html)에서 업무 프로그램을 부르는 연결 코드.
- 브라우저 안의 Python(Pyodide)에서 실행됩니다. 파일은 서버로 올라가지 않습니다.
- 약관에서 추출:        benefit_extractor2.run()   그대로 사용
- 개정 후 자동화:       stat_builder.build()       그대로 사용
- 결과 엑셀을 화면 미리보기용 JSON 으로도 돌려줍니다 (엑셀의 행 색 표시 포함).
"""
import json
import math

import openpyxl

import benefit_extractor2 as be
import stat_builder as sb

PREVIEW_MAX_ROWS = 2000      # 미리보기에 보여 줄 최대 행 수 (엑셀 파일에는 전부 들어 있음)


def _cell(v):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if hasattr(v, "item"):          # numpy 숫자 → 파이썬 숫자
        v = v.item()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def _fill(cell):
    """셀 배경색 → '#RRGGBB' (색 없으면 None)"""
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return None
    rgb = f.fgColor.rgb if f.fgColor is not None else None
    if not isinstance(rgb, str) or len(rgb) < 6:
        return None
    return "#" + rgb[-6:].upper()


def _sheet(ws, hdr_row=1, cols=None):
    """시트 → {name, columns, rows, colors, total}. cols: 보여 줄 열 번호 목록 (None 이면 머리글 있는 열 전부)"""
    if cols is None:
        cols = [c for c in range(1, ws.max_column + 1) if ws.cell(hdr_row, c).value not in (None, "")]
    rows, colors = [], []
    total = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        vals = [_cell(ws.cell(r, c).value) for c in cols]
        if all(v in (None, "") for v in vals):
            continue
        total += 1
        if len(rows) < PREVIEW_MAX_ROWS:
            rows.append(vals)
            colors.append(_fill(ws.cell(r, cols[0])))
    return {"name": ws.title, "columns": [str(ws.cell(hdr_row, c).value) for c in cols],
            "rows": rows, "colors": colors, "total": total}


def sheets_json(xlsx_path):
    """약관 추출 결과 엑셀의 시트들 → 미리보기 JSON (엑셀에 적힌 값 그대로: '01' 은 '01')"""
    wb = openpyxl.load_workbook(xlsx_path)
    return json.dumps([_sheet(ws) for ws in wb.worksheets], ensure_ascii=False)


def run_job(src, out_dir, code_xlsx=None, log=print):
    """[약관에서 추출] 실행: 추출 → 엑셀 저장. (엑셀 경로, 미리보기 JSON)"""
    path = be.run(src, out_dir, code_xlsx or None, log=log)
    return path, sheets_json(path)


def build_job(in_xlsx, out_dir, log=print):
    """[개정 후 자동화] 실행: 양식 엑셀 → 1번 시트 채움 + 검토 시트. (엑셀 경로, 미리보기 JSON)"""
    path = sb.build(in_xlsx, out_dir, log)
    wb = openpyxl.load_workbook(path)
    ws_t = wb.worksheets[sb.SHEET_TARGET]
    hdr, _, all_cols = sb.find_header(ws_t, sb.TABLE_COLS, ws_t.title)
    sheets = [_sheet(ws_t, hdr, sorted(all_cols.values()))]
    for name in ("검토_상품코드", "검토_요약"):
        if name in wb.sheetnames:
            sheets.append(_sheet(wb[name]))
    return path, json.dumps(sheets, ensure_ascii=False)
