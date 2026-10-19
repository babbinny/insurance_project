# -*- coding: utf-8 -*-
"""
통계테이블 자동 작성 프로그램 (1단계: 상품코드 변경 이어받기)
----------------------------------------------------------------
입력: 아래 5개 시트가 들어 있는 엑셀 파일 (양식 파일 그대로)
    1번 시트  통계테이블 작성 대상  ← 여기에 결과가 채워짐
    2번 시트  통계테이블 기존      (개정 전 테이블)
    3번 시트  주보종코드-개별상품코드 관계 테이블 (SP_SPCD, SP_SPCD1)
    4번 시트  변경전,후 상품코드 (변경전, 변경후)
    5번 시트  급부코드 체계 (급부코드, 급부명)

처리:
    4번 시트의 '변경후' 개별상품코드만 작성 대상
      → 2번 시트에서 '변경전' 코드의 행들을 가져오고
      → 3번 시트에서 '변경후' 코드(SP_SPCD1)의 주보종코드(SP_SPCD)를 찾아
      → MAIN_PROD_CODE / SP_PROD_CODE 만 새 코드로 바꿔 1번 시트에 채움
    (급부코드·산출기준·지급율·지급주기·적립금·면책·경과년은 기존 값 그대로)

결과: 원본은 건드리지 않고 '통계테이블_자동작성_날짜시간.xlsx' 로 새로 저장
    [1번 시트]     채워진 통계테이블
    [검토_상품코드] 개별상품코드별 처리 상태 (기존 행 없음, 신규 등은 색 표시)
    [검토_요약]     상태별 건수

- 문서보안(DRM)이 걸린 엑셀이나 구버전(.xls) 파일은 PC의 Excel을 화면에 안 보이게 열어서 값만 읽습니다.
  (이때 결과 파일에는 양식의 색·서식이 빠지고 값만 들어갑니다)

Spyder 에서 F5 로 실행하면 창이 뜹니다. (추가 설치 필요 없음)
"""
import os
import queue
import zipfile
import datetime
import threading
from collections import OrderedDict

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.exceptions import InvalidFileException

# =====================================================================
# 설정
# =====================================================================

# 시트 위치 (0 = 첫 번째 시트). 시트 이름이 바뀌어도 순서로 찾습니다.
SHEET_TARGET = 0
SHEET_OLD = 1
SHEET_REL = 2
SHEET_MAP = 3
SHEET_CODE = 4

# 통계테이블 칼럼 (영문 칼럼명이 적힌 행을 찾아서 그 아래부터 데이터로 봄)
TABLE_COLS = ["MAIN_PROD_CODE", "SP_PROD_CODE", "SCRP_RISK_PRM_CODE", "CALC_TYPE1", "PAY_RATE1",
              "PAY_TERM1", "END_RESRV", "MNGR_START_DAYS", "DUR_YR1", "DUR_YR2"]
REL_MAIN, REL_SP = "SP_SPCD", "SP_SPCD1"      # 관계 테이블 칼럼명
MAP_BEFORE, MAP_AFTER = "변경전", "변경후"      # 변경전/후 칼럼명
CODE_COL, CODE_NAME_COL = "급부코드", "급부명"

HEADER_SEARCH_ROWS = 10                       # 위에서 몇 행 안에서 칼럼명을 찾을지

# =====================================================================
# 엑셀 읽기 도우미
# =====================================================================


def norm(v):
    """코드값 정리: 앞뒤 공백 제거, 숫자로 저장된 코드도 문자로"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def find_header(ws, names, sheet_label):
    """칼럼명이 모두 들어 있는 행 찾기 → (행번호, {칼럼명: 열번호}, 그 행의 전체 칼럼)"""
    for r in range(1, min(ws.max_row, HEADER_SEARCH_ROWS) + 1):
        row = {norm(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1) if norm(ws.cell(r, c).value)}
        if all(n in row for n in names):
            return r, {n: row[n] for n in names}, row
    raise RuntimeError(f"[{sheet_label}] 시트에서 칼럼명 {names} 을(를) 찾지 못했습니다. "
                       f"(위에서 {HEADER_SEARCH_ROWS}행 안에 있어야 합니다)")


def read_rows(ws, hdr_row, cols):
    """칼럼명 행 아래 데이터 읽기. 코드 칼럼이 비어 있는 행은 건너뜀"""
    out = []
    first = list(cols.values())[0]
    for r in range(hdr_row + 1, ws.max_row + 1):
        if not norm(ws.cell(r, first).value):
            continue
        out.append({n: ws.cell(r, c).value for n, c in cols.items()})
    return out


def extra_cols(all_cols, used):
    """칼럼명 행에서 사용하지 않은 나머지 칼럼 (상품명/특약명 등) → 검토 시트에 함께 표시"""
    return {n: c for n, c in all_cols.items() if n not in used}


# ---------------------------------------------------------------------
# Excel 프로그램으로 읽기 (문서보안(DRM) 파일, .xls 파일용)
#   파이썬이 직접 못 여는 파일은 PC의 Excel을 화면에 안 보이게 열어서 시트별 값만 읽음
# ---------------------------------------------------------------------


def _com_value(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)              # Excel 숫자 1.0 → 1
    return v


def read_with_excel(path):
    """PC의 Excel로 열어 시트별 값만 읽어 openpyxl 통합문서로 만듦 (시트 순서·이름·셀 위치 그대로, 서식 제외)"""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        raise RuntimeError("일반 엑셀(.xlsx) 형식이 아니고(문서보안/구버전 .xls 의심), win32com 이 없어 Excel로도 읽지 못했습니다.")
    pythoncom.CoInitialize()          # 실행 버튼 작업(별도 스레드)에서 Excel을 쓰려면 필요
    app = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        xb = app.Workbooks.Open(os.path.abspath(path), UpdateLinks=0, ReadOnly=True)
        try:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            for xs in xb.Worksheets:
                ws = wb.create_sheet(title=str(xs.Name)[:31])
                used = xs.UsedRange
                vals = used.Value2            # 날짜 등도 숫자 그대로 (형식 변환 없이)
                if vals is None:
                    continue
                if not isinstance(vals, tuple):
                    vals = ((vals,),)        # 셀 하나만 있는 시트
                r0, c0 = used.Row, used.Column
                for i, row in enumerate(vals):
                    for j, v in enumerate(row):
                        if v is not None and v != "":
                            ws.cell(r0 + i, c0 + j).value = _com_value(v)
        finally:
            xb.Close(False)
        return wb
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def load_input(path, log=print):
    """일반 .xlsx 는 바로 열고, 안 되면(문서보안/.xls) Excel 프로그램으로 값만 읽음"""
    try:
        return openpyxl.load_workbook(path)
    except (zipfile.BadZipFile, InvalidFileException):
        log("  (참고) 일반 엑셀 형식이 아닙니다(문서보안/구버전 .xls 의심). PC의 Excel로 열어 값만 읽습니다.")
    try:
        wb = read_with_excel(path)
    except RuntimeError:
        raise
    except Exception as e:
        raise RuntimeError(f"Excel로 열기 실패: {e}")
    log("  Excel로 읽음 → 결과 파일에는 양식의 색·서식이 빠지고 값만 들어갑니다.")
    return wb


# =====================================================================
# 핵심 처리
# =====================================================================


def build(in_path, out_dir, log=print):
    wb = load_input(in_path, log)
    if len(wb.worksheets) < 4:
        raise RuntimeError("시트가 4개 이상 있어야 합니다 (작성대상, 기존, 관계, 변경전후).")
    ws_t, ws_o = wb.worksheets[SHEET_TARGET], wb.worksheets[SHEET_OLD]
    ws_r, ws_m = wb.worksheets[SHEET_REL], wb.worksheets[SHEET_MAP]
    ws_c = wb.worksheets[SHEET_CODE] if len(wb.worksheets) > SHEET_CODE else None

    t_hdr, t_cols, _ = find_header(ws_t, TABLE_COLS, ws_t.title)
    o_hdr, o_cols, _ = find_header(ws_o, TABLE_COLS, ws_o.title)
    r_hdr, r_cols, r_all = find_header(ws_r, [REL_MAIN, REL_SP], ws_r.title)
    m_hdr, m_cols, m_all = find_header(ws_m, [MAP_BEFORE, MAP_AFTER], ws_m.title)
    r_extra = extra_cols(r_all, [REL_MAIN, REL_SP])

    old_rows = read_rows(ws_o, o_hdr, o_cols)
    rel = []
    for r in range(r_hdr + 1, ws_r.max_row + 1):
        main, sp = norm(ws_r.cell(r, r_cols[REL_MAIN]).value), norm(ws_r.cell(r, r_cols[REL_SP]).value)
        if main and sp:
            rel.append({"main": main, "sp": sp,
                        "extra": {n: ws_r.cell(r, c).value for n, c in r_extra.items()}})
    maps = []                                   # 변경전이 비어 있어도(신규) 변경후가 있으면 읽음
    for r in range(m_hdr + 1, ws_m.max_row + 1):
        row = {n: ws_m.cell(r, c).value for n, c in m_cols.items()}
        if norm(row[MAP_AFTER]):
            maps.append(row)
    log(f"기존 테이블 {len(old_rows)}행 / 관계 {len(rel)}건 / 변경전후 {len(maps)}건")

    code_names = {}
    if ws_c is not None:
        try:
            c_hdr, c_cols, _ = find_header(ws_c, [CODE_COL, CODE_NAME_COL], ws_c.title)
            for row in read_rows(ws_c, c_hdr, c_cols):
                code_names[norm(row[CODE_COL])] = norm(row[CODE_NAME_COL])
        except RuntimeError as e:
            log(f"  (참고) {e}")

    # 기존 테이블: 개별상품코드 → 주보종코드별 행 목록
    old_by_sp = OrderedDict()
    for row in old_rows:
        sp, main = norm(row["SP_PROD_CODE"]), norm(row["MAIN_PROD_CODE"])
        old_by_sp.setdefault(sp, OrderedDict()).setdefault(main, []).append(row)

    # 관계 테이블: 개별상품코드(SP_SPCD1) → 주보종코드(SP_SPCD) 목록 (+ 상품명 등 나머지 칼럼)
    mains_of = OrderedDict()
    rel_extra = {}
    for x in rel:
        lst = mains_of.setdefault(x["sp"], [])
        if x["main"] not in lst:
            lst.append(x["main"])
        rel_extra[(x["main"], x["sp"])] = x["extra"]

    # 작성 대상 = 4번 시트의 '변경후' 코드만 (같은 변경후가 여러 번 나오면 변경전 코드를 모음)
    targets = OrderedDict()
    for m in maps:
        before, after = norm(m[MAP_BEFORE]), norm(m[MAP_AFTER])
        if not after:
            continue
        lst = targets.setdefault(after, [])
        if before and before not in lst:
            lst.append(before)

    out_rows, hold_rows, review = [], [], []
    for after, befores in targets.items():
        note = []
        if len(befores) > 1:
            note.append(f"변경전 코드가 {len(befores)}개 ({', '.join(befores)}) → 행을 합침")

        # 변경전 코드의 기존 행 가져오기
        src, seen = [], set()
        for code in befores:
            by_main = old_by_sp.get(code)
            if not by_main:
                note.append(f"기존 테이블에 {code} 행 없음")
                continue
            if len(by_main) > 1:
                vals = {tuple(tuple(norm(r[c]) for c in TABLE_COLS[2:]) for r in rs) for rs in by_main.values()}
                if len(vals) > 1:
                    note.append(f"{code}: 기존 주보종({', '.join(by_main)})마다 값이 달라 {list(by_main)[0]} 기준")
            for r in list(by_main.values())[0]:
                key = tuple(norm(r[c]) for c in TABLE_COLS[2:])
                if key not in seen:
                    seen.add(key)
                    src.append(r)

        mains = mains_of.get(after, [])
        if len(mains) > 1:
            note.append(f"주보종코드가 {len(mains)}개 ({', '.join(mains)}) → 주보종마다 작성")

        if not befores:
            status = "신규(변경전 없음)"
        elif not src:
            status = "기존 행 없음"
        elif not mains:
            status = "주보종 모름"
        elif note:
            status = "확인필요"
        else:
            status = "완료"

        bens = OrderedDict()
        for r in src:
            c = norm(r["SCRP_RISK_PRM_CODE"])
            bens[c] = code_names.get(c, "")

        for main in (mains or [""]):
            rows = []
            for r in src:
                new = dict(r)
                new["MAIN_PROD_CODE"], new["SP_PROD_CODE"] = main, after
                rows.append(new)
            if main:
                out_rows += rows
            else:
                hold_rows += rows          # 주보종을 몰라 1번 시트에 못 넣는 행 → 보류 시트
            review.append(OrderedDict([
                ("주보종코드", main), ("개별상품코드(변경후)", after),
                ("변경전코드", ", ".join(befores)),
                ("상태", status), ("작성행수", len(rows) if main else 0),
                ("급부코드", ", ".join(bens)),
                ("급부명", ", ".join(v or "?" for v in bens.values())),
                ("비고", " / ".join(note) if main or not src else
                 " / ".join(note + [f"관계 테이블에 {after} 가 없어 주보종코드를 모름 → '보류_주보종없음' 시트에 {len(rows)}행"])),
                *[(n, v) for n, v in rel_extra.get((main, after), {}).items()],
            ]))

    # ---- 1번 시트에 쓰기 (칼럼명 행 아래 기존 내용은 지우고 새로 채움)
    for r in range(t_hdr + 1, ws_t.max_row + 1):
        for c in t_cols.values():
            ws_t.cell(r, c).value = None
    for i, row in enumerate(out_rows):
        r = t_hdr + 1 + i
        for name, c in t_cols.items():
            v = row[name]
            cell = ws_t.cell(r, c)
            if name in ("MAIN_PROD_CODE", "SP_PROD_CODE", "SCRP_RISK_PRM_CODE", "CALC_TYPE1"):
                cell.value = norm(v)                  # 코드는 문자로 (01 의 0 유지)
                cell.number_format = "@"
            else:
                cell.value = v

    # ---- 검토 시트
    for name in ("검토_상품코드", "검토_요약", "보류_주보종없음"):
        if name in wb.sheetnames:
            del wb[name]
    ws_v = wb.create_sheet("검토_상품코드")
    heads = []
    for rv in review:
        for k in rv:
            if k not in heads:
                heads.append(k)
    ws_v.append(heads)
    colors = {"완료": None, "확인필요": "FFF2CC", "기존 행 없음": "F8CBAD",
              "신규(변경전 없음)": "DDEBF7", "주보종 모름": "F8CBAD"}
    for rv in review:
        ws_v.append([rv.get(h, "") for h in heads])
        color = colors.get(rv["상태"])
        if color:
            for c in range(1, len(heads) + 1):
                ws_v.cell(ws_v.max_row, c).fill = PatternFill("solid", fgColor=color)
    for c in range(1, len(heads) + 1):
        ws_v.cell(1, c).font = Font(bold=True)
        width = max(len(str(ws_v.cell(r, c).value or "")) for r in range(1, ws_v.max_row + 1))
        ws_v.column_dimensions[ws_v.cell(1, c).column_letter].width = min(max(10, width * 1.6), 60)
    ws_v.freeze_panes = "A2"
    ws_v.auto_filter.ref = ws_v.dimensions

    ws_s = wb.create_sheet("검토_요약")
    ws_s.append(["상태", "건수", "설명"])
    desc = {"완료": "변경전 코드의 기존 행을 그대로 이어받음",
            "확인필요": "이어받았지만 비고 확인 필요 (변경전 코드 여러 개, 주보종 여러 개 등)",
            "기존 행 없음": "변경전 코드가 기존 테이블에 없음 → 약관 보고 직접 작성",
            "신규(변경전 없음)": "변경전 코드가 비어 있음 → 신규 상품, 약관 보고 직접 작성",
            "주보종 모름": "변경후 코드가 관계 테이블에 없음 → 행은 '보류_주보종없음' 시트에 있음"}
    counts = OrderedDict((k, 0) for k in desc)
    for rv in review:
        counts[rv["상태"]] = counts.get(rv["상태"], 0) + 1
    for k, n in counts.items():
        ws_s.append([k, n, desc.get(k, "")])
    ws_s.append([])
    ws_s.append(["작성 대상(변경후 코드)", len(targets), "변경전,후 상품코드 시트 기준"])
    ws_s.append(["작성된 행 수", len(out_rows), f"'{ws_t.title}' 시트"])
    if hold_rows:
        ws_s.append(["보류 행 수", len(hold_rows), "'보류_주보종없음' 시트 (주보종코드 채워서 옮기기)"])
        ws_h = wb.create_sheet("보류_주보종없음")
        ws_h.append(TABLE_COLS)
        for row in hold_rows:
            ws_h.append([norm(row[c]) if c in ("MAIN_PROD_CODE", "SP_PROD_CODE", "SCRP_RISK_PRM_CODE", "CALC_TYPE1")
                         else row[c] for c in TABLE_COLS])
        for c in range(1, len(TABLE_COLS) + 1):
            ws_h.cell(1, c).font = Font(bold=True)
            ws_h.column_dimensions[ws_h.cell(1, c).column_letter].width = 20
    for c, w in zip("ABC", (20, 10, 70)):
        ws_s.column_dimensions[c].width = w
    for c in range(1, 4):
        ws_s.cell(1, c).font = Font(bold=True)

    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(out_dir, f"통계테이블_자동작성_{stamp}.xlsx")
    n = 2
    while os.path.exists(out):
        out = os.path.join(out_dir, f"통계테이블_자동작성_{stamp}_{n}.xlsx")
        n += 1
    wb.save(out)

    for k, v in counts.items():
        if v:
            log(f"  {k}: {v}건")
    log(f"작성 대상(변경후 코드): {len(targets)}개 → 작성 행 수: {len(out_rows)}")
    if hold_rows:
        log(f"보류(주보종 모름): {len(hold_rows)}행")
    log(f"결과 파일: {out}")
    return out


# =====================================================================
# 화면
# =====================================================================

def open_file(p):
    try:
        os.startfile(p)
    except Exception:
        pass


def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("통계테이블 자동 작성 (1단계: 상품코드 이어받기)")
    root.geometry("760x420")
    root.lift()
    root.attributes("-topmost", True)
    root.after(800, lambda: root.attributes("-topmost", False))
    v_in, v_out = tk.StringVar(), tk.StringVar()
    q = queue.Queue()

    def pick_in():
        f = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm *.xls")])
        if f:
            v_in.set(f)
            if not v_out.get():
                v_out.set(os.path.dirname(f))

    def pick_out():
        d = filedialog.askdirectory()
        if d:
            v_out.set(d)

    for r, (label, var, cmd) in enumerate([("입력 엑셀(양식)", v_in, pick_in), ("결과 저장 위치", v_out, pick_out)]):
        tk.Label(root, text=label, width=14, anchor="w").grid(row=r, column=0, padx=8, pady=6, sticky="w")
        tk.Entry(root, textvariable=var, width=66).grid(row=r, column=1, padx=4, sticky="ew")
        tk.Button(root, text="찾아보기", command=cmd).grid(row=r, column=2, padx=8)

    logbox = scrolledtext.ScrolledText(root, height=14)
    logbox.grid(row=3, column=0, columnspan=3, padx=8, pady=8, sticky="nsew")
    root.grid_rowconfigure(3, weight=1)
    root.grid_columnconfigure(1, weight=1)

    def log(m):
        q.put(("log", m))

    def poll():
        try:
            while True:
                kind, m = q.get_nowait()
                if kind == "log":
                    logbox.insert("end", m + "\n")
                    logbox.see("end")
                elif kind == "done":
                    btn.config(state="normal")
                    if messagebox.askyesno("완료", "작성이 끝났습니다. 결과 엑셀을 열까요?"):
                        open_file(m)
                elif kind == "error":
                    btn.config(state="normal")
                    messagebox.showerror("오류", m)
        except queue.Empty:
            pass
        root.after(200, poll)

    def start():
        src, out = v_in.get().strip(), v_out.get().strip()
        if not os.path.isfile(src):
            messagebox.showwarning("확인", "입력 엑셀 파일을 선택해 주세요.")
            return
        if not os.path.isdir(out):
            messagebox.showwarning("확인", "결과 저장 위치를 선택해 주세요.")
            return
        btn.config(state="disabled")
        logbox.delete("1.0", "end")
        log("작성 시작...")

        def work():
            try:
                q.put(("done", build(src, out, log)))
            except PermissionError:
                q.put(("error", "파일이 엑셀에서 열려 있으면 닫고 다시 실행해 주세요."))
            except Exception as e:
                q.put(("error", str(e)))

        threading.Thread(target=work, daemon=True).start()

    btn = tk.Button(root, text="실행", width=16, command=start)
    btn.grid(row=2, column=0, columnspan=3, pady=6)
    poll()
    root.mainloop()


if __name__ == "__main__":
    run_gui()
