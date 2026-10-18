# -*- coding: utf-8 -*-
"""
약관 지급기준표 추출 프로그램 (통계테이블 초안 만들기)
----------------------------------------------------------------
- Word 약관(.docx)에서 '급부명 / 지급사유 / 지급금액' 표를 찾아
  급부명, 산출기준, 지급율, 지급주기, 적립금 지급여부, 면책기간, 경과년1/경과년2 를 뽑습니다.
- 감액(예: 1년미만 50%)이 있으면 경과 구간을 자동으로 나눕니다.
    예) 0~12 : 0.5 / 13~999 : 1
- 한 파일에 여러 특약이 들어 있는 책자형 약관(PDF)도 특약 제목을 찾아 특약별로 구분합니다.
- 결과: 통계테이블초안.xlsx
    [초안]        통계테이블 형식 (확인필요 행은 노란색)
    [특약별요약]  특약마다 찾은 급부 수와 급부 목록
    [급부원문]    표에서 읽은 원문 + 근거
    [처리결과]    파일별 성공/실패
- 일반 .docx 는 직접 읽고, 문서보안(DRM) 파일이나 .doc 파일은
  PC의 Word를 화면에 안 보이게 열어서 읽습니다. (추가 설치 필요 없음)
- PDF 도 읽을 수 있습니다 (공개 약관 테스트용). PDF 는 pdfplumber 가 필요합니다.
    Anaconda Prompt 에서: pip install pdfplumber
- Spyder에서 F5로 실행하면 창이 뜹니다.
"""
import os
import re
import queue
import zipfile
import difflib
import datetime
import unicodedata
import threading
import xml.etree.ElementTree as ET

import pandas as pd

# =====================================================================
# 설정 - 회사 기준에 맞게 여기만 고치면 됩니다
# =====================================================================

# 지급기준표의 머리글 판정: 한 행에 아래 단어가 모두(공백 무시) 들어 있으면 머리글로 봄
HEADER_BENEFIT = ["급부명", "급부명칭", "보험금종류", "급부"]      # 급부명 칸 이름 후보
HEADER_REASON = ["지급사유"]                                      # 지급사유 칸 이름 후보
HEADER_AMOUNT = ["지급금액", "지급액"]                             # 지급금액 칸 이름 후보

# 산출기준 코드표 (회사 산출기준)
CALC_BASIS_CODES = {
    "01": "보험가입금액 기준",
    "02": "보험료 (거치형 일시납, 기본보험료) 기준",
    "03": "기납입보험료",
    "10": "MAX(보험가입금액, 기납입보험료 합계) * 추가설계 사망보험금 비율",
    "95": "1구좌(보험료 천만원기준)",
    "96": "1구좌(보험료 1억원기준)",
    "97": "1구좌(보험료 5천만원 기준)",
    "98": "1구좌 (보험료 백만원기준)",
}

# 산출기준 판정: 지급금액 문구(공백 제거)에서 위에서부터 먼저 맞는 규칙의 코드
#   (정규식, 코드). 코드를 모르는 것은 "확인" → '확인필요'로 표시
#   "구좌" 는 1구좌 보험료(천만원/1억원/5천만원/백만원)를 찾아 95~98 로 정함 (UNIT_PREMIUM_CODES)
CALC_BASIS_RULES = [
    (r"(?i)max\(", "10"),                                          # MAX(보험가입금액, 기납입보험료)
    (r"가입금액.{0,20}기납입보험료.{0,20}(?:중|큰|많은|최대)", "10"),  # 가입금액과 기납입보험료 중 큰 금액
    (r"기납입보험료.{0,20}가입금액.{0,20}(?:중|큰|많은|최대)", "10"),
    (r"추가설계", "10"),                                            # 추가설계 사망보험금 비율
    (r"구좌", "구좌"),                                              # 1구좌당 …
    (r"기납입보험료|이미납입한보험료|납입한보험료|총납입보험료", "03"),
    (r"월납|연납|3월납|6월납|일시납|기본보험료|거치", "02"),          # [월납] 월납보험료, 일시납보험료, 기본보험료
    (r"납입보험료", "03"),
    (r"계약자적립액", "확인"),
    (r"보험가입금액|가입금액", "01"),
    (r"보험료(?!.{0,6}면제)", "02"),                                # 그 밖의 '보험료' (납입면제 문구는 제외)
]
# 1구좌 보험료(원) → 산출기준 코드
UNIT_PREMIUM_CODES = {10 ** 7: "95", 10 ** 8: "96", 5 * 10 ** 7: "97", 10 ** 6: "98"}

# 지급금액이 '5,000만원'처럼 금액으로 적힌 표: 표 위 "[기준: 보험가입금액 5,000만원]"으로 나눠 지급율 계산
#   예) 기준 5,000만원 → 지급금액 5,000만원 = 1, 500만원 = 0.1
BASE_AMOUNT_WORDS = ["보험가입금액", "가입금액"]     # '기준' 문구에서 기준금액 앞에 오는 말
MONEY_BASIS_CODE = "01"                          # 금액으로 적힌 지급금액의 산출기준

# 면책기간: "OO보장개시일" 에서 일수를 문서에서 못 찾았을 때 쓰는 기본값
DEFAULT_WAIT_DAYS = {"암": 90}
# 그냥 "보장개시일" (앞에 질병명 없음) 은 면책 없음으로 봄
PLAIN_START_DAYS = 0
# 면책기간이 '년/개월'로 적혀 있을 때 일수로 바꾸는 기준
YEAR_DAYS = 365      # 1년 = 365일
MONTH_DAYS = 30      # 1개월 = 30일 (개월로 적힌 경우는 '확인필요'로 표시)

# 적립금 지급여부 규칙
ONCE_PATTERNS = [r"최초\s*1\s*회"]                                          # → 1
REPEAT_PATTERNS = [r"연\s*간?\s*1\s*회", r"매\s*년", r"매\s*월", r"(?:회|횟)수\s*제한\s*없",
                   r"1\s*일\s*당", r"1\s*회\s*당"]                          # → 0
DEATH_WORDS = ["사망"]                                                        # 급부명에 있으면 → 1
# 특정 시점에 한 번 지급하는 급부(중도급부금·만기 등) → 1  ※ 매월/매년 같은 반복 문구가 있으면 반복(0)이 우선
ONE_TIME_PATTERNS = [r"경과\s*시점", r"종료\s*시점", r"계약\s*해당일에\s*살아\s*있", r"만기\s*(?:시|일)",
                     r"일시\s*(?:금|자금|지급)"]

# 지급금액에 납입주기가 적혀 있으면 [급부원문] 시트 '납입주기' 칸에 표시 (월납/연납 행 구분 참고용)
PAY_CYCLE_WORDS = ["월납", "3월납", "6월납", "연납", "일시납"]

# 지급주기 (통계테이블 PAY_TERM1)
#   급부 발생 때마다 한 번 지급 → "100",  정해진 횟수만큼 매월 지급(예: n회 월지급) → "211",  그 밖 → "100"
PAY_TERM_DEFAULT = "100"
PAY_TERM_MONTHLY_N = "211"
PAY_TERM_MONTHLY_N_PATTERNS = [            # 지급사유+지급금액(공백 제거)에서 찾음. '매월' 같은 월 지급 표현 + 횟수/기간이 함께 있어야 함
    r"매월.{0,30}?\d+회",                   # 매월 계약해당일에 … (120회 한도) / 매월 10만원씩 36회
    r"\d+회.{0,10}?(?:월지급|매월|월분할|매달)",   # 36회 월지급 / 60회에 걸쳐 매월
    r"\d+개월(?:간|동안).{0,20}?(?:매월|매달|월지급)",   # 36개월간 매월
    r"(?:매월|매달).{0,30}?\d+개월(?:간|동안)",         # 매월 … 36개월간
]

# 경과 구간 (월 단위)
END_MONTH = 999

# 특약(상품) 제목 판정: 아래 단어로 끝나는 짧은 줄을 특약 제목으로 봄 (끝의 '약관'은 떼고 판단)
#   예) "1-2. 무배당 갑상선암수술후호르몬약물치료특약", "OO특별약관 약관"
#   한 파일에 여러 특약이 들어 있는 책자형 약관(PDF)에서 표가 어느 특약 것인지 구분할 때 사용
PRODUCT_TITLE_ENDINGS = ["특별약관", "특약", "보통약관"]
# 번호가 붙은 상품 약관 제목도 특약 제목으로 봄. 예) "01. (무)Top3건강보험Ⅱ_2대질병보장형 약관"
PRODUCT_TITLE_PATTERNS = [r"^\d{1,2}\.\s*\S.*(?:보험|특약|특별약관|보장형)\S*\s*약관$"]
PRODUCT_TITLE_MAXLEN = 70
# 문장으로 보이는 줄은 제목에서 제외
PRODUCT_TITLE_EXCLUDE = ["이 특약", "이특약", "합니다", "됩니다", "경우", "에 따라", "하여", "에서"]

# 급부코드 매핑 엑셀(선택): 첫 번째 열 = 급부코드, 두 번째 열 = 급부코드명
# 급부명 비교 시 끝에서 떼어낼 표현 (예: 치료비 ↔ 치료)
NAME_SUFFIXES = ["보험금", "급여금", "지원금", "자금", "비", "금"]

# 급부코드표: 칼럼명으로 찾음 (모든 시트, 위에서 15행 안에서 칼럼명 행을 찾음)
CODE_ID_COLS = ["CD_ID", "급부코드", "SCRP_RISK_PRM_CODE"]
CODE_NM_COLS = ["CD_NM", "급부코드명", "급부명"]

# 급부 매칭
#   1) 급부명 전체가 급부코드명과 거의 같으면(NAME_DIRECT_SCORE 이상) 그 코드 하나
#   2) 아니면 지급사유에 나온 질병명(약관의 "‘OO’이라 함은" 정의 + 급부코드명에 있는 질병명)마다
#      '질병명 + 행위(진단/수술…)' 로 급부코드를 찾음 → 질병명이 여러 개면 급부코드도 여러 개(행도 여러 개)
#   3) 그래도 없으면 급부명과 가장 비슷한 코드 (확인필요)
NAME_DIRECT_SCORE = 0.88
MATCH_OK_SCORE = 0.85        # 이 점수 이상이면 정상, 미만이면 '확인필요'
MATCH_MIN_SCORE = 0.60       # 이 점수 미만이면 코드를 넣지 않음
# 약관 용어 → 급부코드표 용어 (예: 약관의 '암' = 급부코드표의 '일반암')
TERM_ALIAS = {"암": "일반암"}
# 급부의 '행위' (급부명에서 먼저 찾고, 없으면 지급사유에서 찾음)
ACTION_WORDS = ["진단", "수술", "입원", "통원", "치료", "사망", "장해", "항암", "방사선", "약물", "검사",
                "간병", "요양", "골절", "화상", "생존", "중도", "만기"]
# 질병명으로 볼 용어의 끝 글자 (약관 정의에서 '청약일로부터…' 같은 일반 용어를 거르기 위함)
DISEASE_ENDINGS = ["암", "종양", "증", "혈", "질환", "질병", "병", "치매", "경색", "출혈", "골절", "화상", "장해", "결석"]

# =====================================================================
# Word(.docx) 읽기 - 추가 설치 없이 zip/XML 로 직접 읽음
# =====================================================================

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _para_text(p):
    parts = []
    for node in p.iter():
        if node.tag == W + "t" and node.text:
            parts.append(node.text)
        elif node.tag in (W + "tab",):
            parts.append(" ")
        elif node.tag in (W + "br", W + "cr"):
            parts.append(" ")
    return "".join(parts)


def _cell_text(tc):
    paras = [_para_text(p) for p in tc.findall(W + "p")]
    # 셀 안에 표가 또 있으면 그 글자도 포함
    for inner in tc.findall(W + "tbl"):
        for row in _table_rows(inner):
            paras.append(" ".join(row))
    return clean(" ".join(paras))


def _table_rows(tbl):
    """표를 2차원 글자 목록으로. 가로/세로 병합된 칸은 같은 값으로 채움."""
    rows, prev = [], []
    for tr in tbl.findall(W + "tr"):
        row = []
        for tc in tr.findall(W + "tc"):
            tcpr = tc.find(W + "tcPr")
            span, vcont = 1, False
            if tcpr is not None:
                gs = tcpr.find(W + "gridSpan")
                if gs is not None:
                    span = int(gs.get(W + "val", "1"))
                vm = tcpr.find(W + "vMerge")
                if vm is not None and vm.get(W + "val", "continue") == "continue":
                    vcont = True
            col = len(row)
            text = prev[col] if (vcont and col < len(prev)) else _cell_text(tc)
            row.extend([text] * span)
        rows.append(row)
        prev = row
    return rows


def read_docx(path):
    """문서 순서대로 [('p', 글자) 또는 ('tbl', 2차원표)] 목록 반환."""
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(W + "body")
    blocks = []
    for child in body:
        if child.tag == W + "p":
            t = clean(_para_text(child))
            if t:
                blocks.append(("p", t))
        elif child.tag == W + "tbl":
            blocks.append(("tbl", _table_rows(child)))
    return blocks


# ---------------------------------------------------------------------
# Word 프로그램으로 읽기 (문서보안(DRM) 파일, .doc 파일용)
#   파이썬이 직접 못 여는 파일은 PC의 Word를 화면에 안 보이게 열어서 읽음
# ---------------------------------------------------------------------

def _com_clean(s):
    return clean((s or "").replace("\r", " ").replace("\x07", " ").replace("\x0b", " "))


def _com_table_rows(tbl):
    """Word 표 → 2차원 글자 목록. 세로 병합으로 빠진 칸은 윗행 값으로 채움."""
    grid = {}
    cells = tbl.Range.Cells
    for k in range(1, cells.Count + 1):
        c = cells(k)
        grid.setdefault(c.RowIndex, {})[c.ColumnIndex] = _com_clean(c.Range.Text)
    rows, prev = [], []
    for ri in sorted(grid):
        cols = grid[ri]
        texts = [cols[ci] for ci in sorted(cols)]
        width = max(len(prev), max(cols))
        if prev and sorted(cols) == list(range(1, len(cols) + 1)) and len(cols) < width:
            row = prev[: width - len(cols)] + texts          # 앞쪽 칸(급부명 등)이 병합된 경우
        else:
            row = [cols.get(ci, prev[ci - 1] if ci - 1 < len(prev) else "") for ci in range(1, width + 1)]
        rows.append(row)
        prev = row
    return rows


def _add_paras(blocks, text):
    for line in (text or "").split("\r"):
        t = _com_clean(line)
        if t:
            blocks.append(("p", t))


class WordReader:
    def __init__(self):
        self.app = None

    def read(self, path):
        if self.app is None:
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()      # 실행 버튼 작업(별도 스레드)에서 Word를 쓰려면 필요
            self.app = win32com.client.DispatchEx("Word.Application")
            self.app.Visible = False
            self.app.DisplayAlerts = 0
        doc = self.app.Documents.Open(os.path.abspath(path), ConfirmConversions=False,
                                      ReadOnly=True, AddToRecentFiles=False, Visible=False)
        try:
            tables = [doc.Tables(i) for i in range(1, doc.Tables.Count + 1)]
            tables.sort(key=lambda t: t.Range.Start)
            blocks, pos = [], 0
            for t in tables:
                s, e = t.Range.Start, t.Range.End
                if s < pos:
                    continue
                _add_paras(blocks, doc.Range(pos, s).Text)
                blocks.append(("tbl", _com_table_rows(t)))
                pos = e
            _add_paras(blocks, doc.Range(pos, doc.Content.End).Text)
            return blocks
        finally:
            doc.Close(False)

    def quit(self):
        if self.app is not None:
            try:
                self.app.Quit()
            except Exception:
                pass
            self.app = None
            try:
                import pythoncom
                pythoncom.CoUninitialize()
            except Exception:
                pass


# ---------------------------------------------------------------------
# PDF 읽기 (공개 약관 PDF 테스트용)  ※ pdfplumber 필요: pip install pdfplumber
#   지급기준표는 PDF 프로그램의 자동 표 인식 대신 직접 칸을 나눔:
#   1) '급부명 / 지급사유 / 지급금액' 머리글을 찾고
#   2) 머리글 위치로 세 칸의 경계를 정하고
#   3) '지급사유' 칸을 가로지르는 가로선을 행 경계로 봄
#      (급부명 칸의 강조 상자, 지급금액 칸만 나누는 선(월납/연납)은 행을 나누지 않음)
# ---------------------------------------------------------------------

ROW_LINE_COVER = 0.7      # 가로선이 칸 너비의 70% 이상 덮으면 그 칸의 경계선으로 봄


def _is_benefit_header(row):
    row = [c or "" for c in row]
    return (find_col(row, HEADER_BENEFIT) is not None and find_col(row, HEADER_REASON) is not None
            and find_col(row, HEADER_AMOUNT) is not None)


def _merge_close(xs, tol=3):
    out = []
    for x in sorted(xs):
        if not out or x - out[-1] > tol:
            out.append(x)
    return out


def _hsegs(page):
    """가로선 조각들 [(y, x0, x1)] : 선, 얇은 사각형, 테두리 있는 사각형의 위·아래 변"""
    segs = []
    for ln in page.lines:
        if abs(ln["top"] - ln["bottom"]) < 1.5:
            segs.append(((ln["top"] + ln["bottom"]) / 2, min(ln["x0"], ln["x1"]), max(ln["x0"], ln["x1"])))
    for r in page.rects:
        h = r["bottom"] - r["top"]
        if h <= 2:                                   # 얇은 사각형 = 선
            segs.append(((r["top"] + r["bottom"]) / 2, r["x0"], r["x1"]))
        elif r.get("stroke"):                        # 테두리 있는 칸
            segs.append((r["top"], r["x0"], r["x1"]))
            segs.append((r["bottom"], r["x0"], r["x1"]))
    return segs


def _cover(segs, y, x0, x1, tol=1.5):
    """y 높이의 가로선들이 x0~x1 을 얼마나 덮는지 (0~1)"""
    iv = sorted((max(a, x0), min(b, x1)) for yy, a, b in segs if abs(yy - y) <= tol and b > x0 and a < x1)
    total, cur_a, cur_b = 0.0, None, None
    for a, b in iv:
        if cur_b is None or a > cur_b + 3.5:          # 칸 사이 좁은 틈(3pt)은 이어진 것으로 봄
            if cur_b is not None:
                total += cur_b - cur_a
            cur_a, cur_b = a, b
        else:
            cur_b = max(cur_b, b)
    if cur_b is not None:
        total += cur_b - cur_a
    return total / max(x1 - x0, 1)


def _line_ys(segs, cols, col_i, y0, y1):
    """y0~y1 사이에서 col_i 번째 칸을 가로지르는 가로선의 y 목록"""
    a, b = cols[col_i], cols[col_i + 1]
    ys = _merge_close([y for y, _, _ in segs if y0 <= y <= y1], 1.5)
    return [y for y in ys if _cover(segs, y, a, b) >= ROW_LINE_COVER]


def _word_lines(words):
    lines = {}
    for w in words:
        lines.setdefault(round(w["top"] / 3), []).append(w)
    return [sorted(lines[k], key=lambda w: w["x0"]) for k in sorted(lines)]


def _find_headers(words):
    """머리글 줄 찾기 → [(머리글 top, bottom, [급부명·지급사유·지급금액 글자 중심 x])]"""
    found = []
    for line in _word_lines(words):
        joined, owner = "", []
        for wi, w in enumerate(line):
            t = nospace(clean(w["text"]))
            joined += t
            owner += [wi] * len(t)
        centers = []
        for cands in (HEADER_BENEFIT, HEADER_REASON, HEADER_AMOUNT):
            pos = next((joined.find(c) for c in cands if c in joined), -1)
            if pos < 0:
                break
            w = line[owner[pos]]
            centers.append((w["x0"] + w["x1"]) / 2)
        if len(centers) == 3 and centers == sorted(centers):
            found.append((min(w["top"] for w in line), max(w["bottom"] for w in line), centers))
    return found


def _cols_from_cells(page, htop, hbot, centers):
    """PDF 자동 표 인식 결과에서 머리글 칸 위치를 얻어 세 칸 경계 계산 (칸 사이 틈은 반으로 나눔)"""
    try:
        tables = page.find_tables()
    except Exception:
        return None
    for t in tables:
        for row in t.rows:
            cells = [c for c in row.cells if c and c[1] <= hbot and c[3] >= htop]
            if not cells:
                continue
            pick = []
            for cx in centers:
                inside = [c for c in cells if c[0] - 1 <= cx <= c[2] + 1]
                if not inside:
                    break
                pick.append(min(inside, key=lambda c: c[2] - c[0]))
            if len(pick) == 3 and pick[0][2] <= pick[1][0] + 1 and pick[1][2] <= pick[2][0] + 1:
                left, right = min(c[0] for c in cells), max(c[2] for c in cells)
                return [left, (pick[0][2] + pick[1][0]) / 2, (pick[1][2] + pick[2][0]) / 2, right]
    return None


def _cols_from_gaps(page, segs, words, htop, hbot, centers):
    """선/칸 정보가 부족할 때: 머리글 아래 글자들 사이 빈 공간으로 칸 경계 추정"""
    # 표 좌우 끝: 머리글 근처 가로선의 범위
    near = [(a, b) for y, a, b in segs if htop - 30 <= y <= hbot + 30 and a <= centers[0] and b >= centers[0]]
    near += [(a, b) for y, a, b in segs if htop - 30 <= y <= hbot + 30 and a <= centers[2] and b >= centers[2]]
    if near:
        left, right = min(a for a, _ in near), max(b for _, b in near)
    else:
        left, right = centers[0] - 60, centers[2] + 80
    body = [w for w in words if hbot < w["top"] < hbot + 250 and left <= w["x0"] <= right]
    # 일반 문단 줄(1·2번째 칸에 걸쳐 이어지는 글) 이후는 제외
    for line in _word_lines(body):
        x0, prev = line[0]["x0"], line[0]["x1"]
        para = False
        for w in line[1:] + [None]:
            if w is None or w["x0"] - prev > 8:
                if x0 <= centers[0] + 15 and prev >= centers[1]:
                    para = True
                    break
                if w is not None:
                    x0 = w["x0"]
            if w is not None:
                prev = w["x1"]
        if para:
            lim = min(w["top"] for w in line)
            body = [w for w in body if w["top"] < lim]
            break
    used = [False] * (int(right) + 2)
    for w in body:
        for x in range(int(w["x0"]), int(w["x1"]) + 1):
            if 0 <= x < len(used):
                used[x] = True
    cuts = []
    for a, b in ((centers[0], centers[1]), (centers[1], centers[2])):
        best, start = (0, (a + b) / 2), None
        for x in range(int(a), int(b) + 1):
            if x < len(used) and not used[x]:
                start = x if start is None else start
                if x - start + 1 > best[0]:
                    best = (x - start + 1, start + (x - start) / 2)
            else:
                start = None
        cuts.append(best[1])
    return [left, cuts[0], cuts[1], right]


def _crosses(words, cols, y0, y1):
    """y0~y1 사이에 칸 경계를 가로지르는 글자가 있으면(=표 밖 일반 문단) True"""
    for w in words:
        if y0 < (w["top"] + w["bottom"]) / 2 < y1 and any(w["x0"] + 1 < c < w["x1"] - 1 for c in cols[1:-1]):
            return True
    return False


def _table_extent(segs, words, cols, start_y, stop_y):
    """start_y 부터 '지급사유' 칸을 가로지르는 가로선이 이어지는 범위 → 행 경계 목록"""
    ys = _line_ys(segs, cols, 1, start_y, stop_y)
    if not ys:
        return []
    run = [ys[0]]
    for y in ys[1:]:
        if y - run[-1] > 400 or _crosses(words, cols, run[-1], y):
            break
        run.append(y)
    return run


def _pdf_cell_text(words, x0, x1, y0, y1):
    ws = [w for w in words if x0 <= (w["x0"] + w["x1"]) / 2 < x1 and y0 < (w["top"] + w["bottom"]) / 2 < y1]
    return clean(" ".join(" ".join(w["text"] for w in ln) for ln in _word_lines(ws)))


def _grid_rows(segs, words, cols, ys):
    """행 경계(ys)와 칸 경계(cols)로 [급부명, 지급사유, 지급금액] 행 만들기.
    지급금액 칸만 나누는 선이 있으면(월납/연납) 같은 급부명·지급사유로 행을 나눔"""
    rows = []
    for r0, r1 in zip(ys, ys[1:]):
        name = _pdf_cell_text(words, cols[0], cols[1], r0, r1)
        reason = _pdf_cell_text(words, cols[1], cols[2], r0, r1)
        subs = [y for y in _line_ys(segs, cols, 2, r0 + 2, r1 - 2)]
        bounds = [r0] + subs + [r1]
        amts = [_pdf_cell_text(words, cols[2], cols[3], a, b) for a, b in zip(bounds, bounds[1:])]
        amts = [a for a in amts if a] or [""]
        if not (name or reason or any(amts)):
            continue
        for a in amts:
            rows.append([name, reason, a])
    return rows


def _pdf_page_tables(page, carry):
    """한 쪽에서 지급기준표 찾기 → ([(top, bottom, 행목록, 머리글여부)], 다음 쪽으로 넘길 칸 경계)"""
    try:
        words = page.extract_words()
    except Exception:
        return [], None
    segs = _hsegs(page)
    heads = _find_headers(words)
    found = []
    for i, (htop, hbot, centers) in enumerate(heads):
        cols = _cols_from_cells(page, htop, hbot, centers) or _cols_from_gaps(page, segs, words, htop, hbot, centers)
        stop = heads[i + 1][0] - 2 if i + 1 < len(heads) else page.height
        ys = _table_extent(segs, words, cols, hbot - 3, stop)
        if len(ys) < 2:
            continue
        rows = [["급부명", "지급사유", "지급금액"]] + _grid_rows(segs, words, cols, ys)
        top = htop - 2
        above = [y for y, a, b in segs if htop - 15 <= y < htop and _cover(segs, y, cols[1], cols[2]) >= ROW_LINE_COVER]
        if above:
            top = min(above) - 1
        found.append((top, ys[-1] + 1, rows, True, cols))
    # 앞 쪽에서 이어지는 표 (이 쪽 맨 위, 머리글 없이)
    if carry:
        first_head = heads[0][0] if heads else page.height
        ys = _table_extent(segs, words, carry, 0, first_head - 2)
        if ys and ys[0] < page.height * 0.4:
            # 쪽이 바뀌며 잘린 칸의 나머지 글자(첫 가로선 위)도 표에 포함
            above = [w for w in words if w["bottom"] <= ys[0] + 1
                     and carry[0] - 2 <= w["x0"] and w["x1"] <= carry[-1] + 2]
            top_y = ys[0]
            for ln in sorted(_word_lines(above), key=lambda l: -min(w["top"] for w in l)):
                lt = min(w["top"] for w in ln)
                if top_y - max(w["bottom"] for w in ln) > 40 or _crosses(ln, carry, lt - 1, lt + 30):
                    break
                top_y = lt - 1
            if top_y < ys[0]:
                ys = [top_y] + ys
            if len(ys) >= 2:
                found.append((ys[0] - 1, ys[-1] + 1, _grid_rows(segs, words, carry, ys), False, carry))
    found.sort(key=lambda f: f[0])
    nxt = None
    if found:
        top, bottom, rows, _, cols = found[-1]
        below = [w for w in words if bottom + 2 < w["top"] < page.height * 0.9]
        if bottom > page.height * 0.6 and not below:     # 표가 쪽 끝까지 이어짐 → 다음 쪽으로 계속될 수 있음
            nxt = cols
    return found, nxt


PDF_PROGRESS_PAGES = 100   # PDF 를 읽을 때 이 쪽수마다 진행 상황을 로그에 표시


def read_pdf(path, log=None):
    try:
        import pdfplumber
    except ImportError:
        raise RuntimeError("PDF를 읽으려면 pdfplumber 가 필요합니다 (Anaconda Prompt에서: pip install pdfplumber)")
    blocks = []
    carry = None
    with pdfplumber.open(path) as pdf:
        n_pages = len(pdf.pages)
        for pno, page in enumerate(pdf.pages, 1):
            tables, carry = _pdf_page_tables(page, carry)
            boxes = [(t[4][0] - 2, t[0], t[4][-1] + 2, t[1]) for t in tables]
            items = [(t[0], "tbl", t[2]) for t in tables if t[2]]

            def outside(obj, boxes=boxes):
                if obj.get("object_type") != "char":
                    return True
                cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
                return not any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes)

            rest = page.filter(outside)
            try:
                lines = rest.extract_text_lines()
                items += [(ln["top"], "p", clean(ln["text"])) for ln in lines if clean(ln["text"])]
            except Exception:
                y = 0
                for ln in (rest.extract_text() or "").split("\n"):
                    if clean(ln):
                        items.append((y, "p", clean(ln)))
                        y += 0.001
            items.sort(key=lambda x: x[0])
            blocks += [(k, d) for _, k, d in items]
            page.close()        # 읽은 쪽의 글자·선 정보를 메모리에서 비움 (큰 PDF에서 메모리 부족 방지)
            if log and n_pages > PDF_PROGRESS_PAGES and (pno % PDF_PROGRESS_PAGES == 0 or pno == n_pages):
                log(f"  {os.path.basename(path)}: {pno}/{n_pages}쪽 읽음")
    text = " ".join(d for k, d in blocks if k == "p")
    if text.count("(cid:") > 20:
        raise RuntimeError("PDF 글자를 읽을 수 없음(글꼴 문제). Word 파일로 시도해 주세요.")
    return blocks


def read_any(path, word_reader, log=None):
    """일반 .docx 는 바로 읽고, 안 되면(DRM/.doc) Word 프로그램으로 읽음. (blocks, 읽은방식)"""
    if path.lower().endswith(".pdf"):
        return read_pdf(path, log), "PDF"
    if path.lower().endswith((".docx", ".docm")):
        try:
            return read_docx(path), "직접"
        except zipfile.BadZipFile:
            pass        # 문서보안(DRM) 등으로 일반 형식이 아님 → Word로 읽기
    try:
        return word_reader.read(path), "Word"
    except ImportError:
        raise RuntimeError("일반 Word 형식이 아니고(DRM 의심), win32com 이 없어 Word로도 읽지 못함")
    except Exception as e:
        raise RuntimeError(f"Word로 열기 실패: {e}")


# =====================================================================
# 규칙
# =====================================================================


def clean(s):
    s = unicodedata.normalize("NFC", s or "")   # PDF에서 한글이 자모로 쪼개져 나오는 경우 대비
    return re.sub(r"\s+", " ", s).strip()


def nospace(s):
    return re.sub(r"\s+", "", s or "")


def find_col(header, candidates):
    for i, h in enumerate(header):
        h2 = nospace(h)
        if any(c in h2 for c in candidates):
            return i
    return None


def split_name(raw):
    """'갑상선암 수술 후 호르몬약물치료비 (제2-2조)' → ('갑상선암 수술 후 호르몬약물치료비', '제2-2조')"""
    m = re.search(r"[\(\[]\s*(제\s*[\d\-\s]+조[^\)\]]*)[\)\]]", raw)
    article = nospace(m.group(1)) if m else ""
    name = clean(re.sub(r"[\(\[]\s*제\s*[\d\-\s]+조[^\)\]]*[\)\]]", "", raw))
    return name, article


MONEY_UNITS = {"억": 10 ** 8, "천만": 10 ** 7, "백만": 10 ** 6, "십만": 10 ** 5, "만": 10 ** 4, "천": 10 ** 3, "": 1}
MONEY_RE = r"((?:\d+(?:\.\d+)?(?:억|천만|백만|십만|만|천)?)+)원"


def find_money(text):
    """'5,000만원', '1억 5,000만원', '500만 원', '100,000원' → 원 단위 숫자 목록"""
    t = nospace(text).replace(",", "")
    out = []
    for m in re.finditer(MONEY_RE, t):
        won = sum(float(n) * MONEY_UNITS[u] for n, u in re.findall(r"(\d+(?:\.\d+)?)(억|천만|백만|십만|만|천)?", m.group(1)))
        if won > 0:
            out.append(won)
    return out


def find_base_amount(texts):
    """'[기준: 보험가입금액 5,000만원]' → 50000000 (못 찾으면 None)"""
    for t in texts:
        s = nospace(t)
        if "기준" not in s:
            continue
        for w in BASE_AMOUNT_WORDS:
            if w in s:
                money = find_money(s[s.index(w) + len(w):])
                if money:
                    return money[0], clean(t)
    return None, ""


def find_unit_premium(texts):
    """'1구좌(보험료 1,000만원 기준)', '[기준: 1구좌 = 보험료 1억원]' → (1구좌 보험료(원), 근거). 못 찾으면 (None, '')"""
    for t in texts:
        s = nospace(t).replace(",", "")
        if "구좌" not in s:
            continue
        s = re.sub(r"(?<![\d.])(억|천만|백만|십만)원", r"1\1원", s)      # '천만원' → '1천만원'
        m = re.search(r"구좌.{0,15}?보험료(.{0,20})", s)
        money = find_money(m.group(1)) if m else []
        if money:
            return money[0], clean(t)
    return None, ""


def _unit_basis(unit_premium):
    code = UNIT_PREMIUM_CODES.get(unit_premium)
    if code:
        return code, f"1구좌(보험료 {int(unit_premium):,}원)"
    if unit_premium:
        return "확인", f"1구좌 보험료 {int(unit_premium):,}원은 산출기준표에 없음"
    return "확인", "1구좌 보험료를 못 찾음"


def get_basis(amount_text, unit_premium=None):
    """지급금액 → (산출기준 코드, 근거 문구). unit_premium: 지급사유·지급금액에 '구좌'가 있을 때 찾은 1구좌 보험료"""
    t = nospace(amount_text)
    for pat, code in CALC_BASIS_RULES:
        m = re.search(pat, t)
        if m:
            return _unit_basis(unit_premium) if code == "구좌" else (code, m.group(0))
    if unit_premium:                         # 지급사유에만 '1구좌'가 있고 지급금액은 금액만 적힌 경우
        return _unit_basis(unit_premium)
    if find_money(amount_text):              # '5,000만원' 처럼 금액만 적힌 경우
        return MONEY_BASIS_CODE, "금액"
    return "확인", ""


# 감액 문구: '1년미만 … 50%'  (get_reduction 과 get_rate 가 함께 씀)
REDUCTION_RE = r"(\d+)\s*년\s*(?:미만|이내)([^%]{0,60}?)(\d+(?:\.\d+)?)\s*%"


def get_rate(amount_text, base=None):
    """지급금액 → (지급율, 확인필요메모)
    감액 문구의 %는 기본 지급율이 아니므로 빼고 봄
      예) '보험가입금액 (다만, 1년미만 진단확정시 보험가입금액의 50% 지급)' → 1 (감액 50%는 경과 구간에서 따로 적용)"""
    amount_text = re.sub(REDUCTION_RE, " ", amount_text)
    pcts = re.findall(r"(\d+(?:\.\d+)?)\s*%", amount_text)
    if len(pcts) == 1:
        return round(float(pcts[0]) / 100, 4), ""
    if len(pcts) > 1:
        return round(float(pcts[0]) / 100, 4), f"지급금액에 %가 여러 개({', '.join(pcts)}%)"
    money = find_money(amount_text)
    if money:                                # 금액 ÷ 기준금액 (보험가입금액 또는 1구좌 보험료)
        if not base:
            return None, "지급금액이 금액인데 [기준: 보험가입금액] 을 못 찾음"
        memo = f"지급금액에 금액이 여러 개({len(money)}개)" if len(money) > 1 else ""
        return round(money[0] / base, 4), memo
    t = nospace(amount_text)
    # "보험가입금액", "총 납입보험료", "월납보험료" 처럼 기준만 적혀 있으면 100% 로 봄
    if any(re.search(p, t) for p, _ in CALC_BASIS_RULES) and not re.search(r"\d", t):
        return 1.0, ""
    return None, "지급금액에 % / 금액이 없음"


def get_reduction(text):
    """감액 문구 → (년수, 비율, 방식, 원문)
    '1년미만 ... 치료비의 50%'       → 기본지급율 × 0.5 (비율)
    '1년미만 ... 보험가입금액의 50%' → 0.5 그대로 (절대값)"""
    m = re.search(REDUCTION_RE, text)
    if not m:
        return None, None, "", ""
    mode = "절대값" if "가입금액" in nospace(m.group(2)) else "비율"
    return int(m.group(1)), float(m.group(3)) / 100, mode, clean(m.group(0))


def get_wait_days(reason, notes, whole, section=""):
    """지급사유의 'OO보장개시일' → 면책일수. (값, 근거, 확인필요메모)"""
    # 띄어쓰기 기준 한 덩어리만 봄: "중 암보장개시일" → "암", "중, 보장개시일" → ""
    prefixes = re.findall(r"([가-힣A-Za-z]*)보장개시일", re.sub(r"보장\s+개시일", "보장개시일", reason))
    if not prefixes:
        return None, "", "지급사유에 보장개시일 문구 없음"
    named = [p[-6:] for p in prefixes if p]
    if not named:
        return PLAIN_START_DAYS, "보장개시일(질병명 없음)", ""
    prefix = named[0]
    key = prefix + "보장개시일"
    after = r"\s*(?:이|가|째|을|이\s*지난|이\s*경과|경과|지난|후|이\s*되는)"

    def sentences(src):
        # 문단(줄) 단위로 먼저 보고, PDF처럼 문장이 줄바꿈으로 잘린 경우를 위해 문장 단위로도 봄
        yield from re.split(r"(?<=[.。])\s+|\n", src)
        yield from re.split(r"(?<=[.。])\s+", src.replace("\n", " "))

    for src_name, src in (("(주)", notes), ("특약본문", section), ("약관전체", whole)):
        for sent in sentences(src):
            if key not in nospace(sent):
                continue
            # 같은 문장에 '2년(730일)'처럼 둘 다 있으면 '일'을 우선 사용
            m = re.search(r"(\d+)\s*일" + after, sent)
            if m:
                return int(m.group(1)), f"{src_name}: {sent[:120]}", ""
            m = re.search(r"(\d+)\s*년" + after, sent)
            if m:
                return int(m.group(1)) * YEAR_DAYS, f"{src_name}: {sent[:120]}", ""
            m = re.search(r"(\d+)\s*개월" + after, sent)
            if m:
                return (int(m.group(1)) * MONTH_DAYS, f"{src_name}: {sent[:120]}",
                        f"{key} {m.group(1)}개월→{int(m.group(1)) * MONTH_DAYS}일 환산 확인")
    if prefix in DEFAULT_WAIT_DAYS:
        return DEFAULT_WAIT_DAYS[prefix], f"{key} 기본값", ""
    return None, key, f"{key} 일수를 문서에서 못 찾음"


def get_once_flag(name, text):
    t = nospace(text)
    if any(w in nospace(name) for w in DEATH_WORDS):
        return 1, "사망 급부"
    once = any(re.search(p, t) for p in ONCE_PATTERNS)
    rep = any(re.search(p, t) for p in REPEAT_PATTERNS)
    if once and not rep:
        return 1, "최초 1회"
    if rep and not once:
        return 0, "반복 지급"
    if once and rep:
        return None, "최초1회/반복 문구가 함께 있음"
    if any(re.search(p, t) for p in ONE_TIME_PATTERNS):
        return 1, "특정 시점 1회 지급"
    return None, "지급횟수 문구 없음"


def get_pay_cycle(amount_text):
    """'[월납] 월납보험료…' → '월납' (없으면 '')"""
    t = nospace(amount_text)
    found = [w for w in sorted(PAY_CYCLE_WORDS, key=len, reverse=True) if w in t]
    return found[0] if found else ""


def get_pay_term(text):
    """지급사유+지급금액 → (지급주기 코드, 근거). n회 월지급이면 '211', 그 밖은 '100'"""
    t = nospace(text)
    for p in PAY_TERM_MONTHLY_N_PATTERNS:
        m = re.search(p, t)
        if m:
            return PAY_TERM_MONTHLY_N, m.group(0)
    return PAY_TERM_DEFAULT, ""


def norm_benefit(s):
    """급부명 비교용: 괄호·공백·로마숫자·끝의 '자금/비/금' 등 제거"""
    s = nospace(re.sub(r"[\(\[【].*?[\)\]】]", "", s or ""))
    s = re.sub(r"(?:Ⅰ|Ⅱ|Ⅲ|Ⅳ|Ⅴ|II|III|IV)+$", "", s)
    for suf in NAME_SUFFIXES:
        if s.endswith(suf) and len(s) > len(suf) + 1:
            s = s[: -len(suf)]
            break
    return s


def actions_in(text):
    t = nospace(text)
    return [a for a in ACTION_WORDS if a in t]


def disease_part(name_n):
    """급부코드명(정리된 것)에서 행위어를 뺀 질병 부분. 예) 기타피부암진단 → 기타피부암"""
    d = name_n
    for a in sorted(ACTION_WORDS, key=len, reverse=True):
        d = d.replace(a, "")
    return d


def _find_col_name(header, cands):
    hs = [nospace(str(h)).upper() if h is not None else "" for h in header]
    for c in cands:
        if c.upper() in hs:
            return hs.index(c.upper())
    return None


def load_code_table(path):
    """급부코드표 읽기: 'CD_ID'(급부코드) / 'CD_NM'(급부명) 칼럼을 이름으로 찾음"""
    if not path:
        return None
    if path.lower().endswith(".csv"):
        try:
            sheets = {"csv": pd.read_csv(path, header=None, dtype=str, encoding="utf-8-sig")}
        except UnicodeDecodeError:
            sheets = {"csv": pd.read_csv(path, header=None, dtype=str, encoding="cp949")}
    else:
        sheets = pd.read_excel(path, header=None, dtype=str, sheet_name=None)
    for sname, raw in sheets.items():
        for hr in range(min(15, len(raw))):
            header = list(raw.iloc[hr])
            ci, ni = _find_col_name(header, CODE_ID_COLS), _find_col_name(header, CODE_NM_COLS)
            if ci is None or ni is None:
                continue
            df = raw.iloc[hr + 1:, [ci, ni]].copy()
            df.columns = ["급부코드", "급부코드명"]
            df["급부코드"] = df["급부코드"].fillna("").map(lambda v: str(v).strip())
            df["급부코드명"] = df["급부코드명"].fillna("").map(lambda v: clean(str(v)))
            df = df[(df["급부코드"] != "") & (df["급부코드명"] != "")].drop_duplicates("급부코드")
            df["_n"] = df["급부코드명"].map(norm_benefit)
            df["_d"] = df["_n"].map(disease_part)
            df["_a"] = df["급부코드명"].map(actions_in)
            return df.reset_index(drop=True)
    raise RuntimeError(f"급부코드표에서 칼럼명 {CODE_ID_COLS[0]} / {CODE_NM_COLS[0]} 을(를) 찾지 못했습니다.")


def defined_terms(whole):
    """약관 전체에서 "‘OO’이라 함은" 으로 정의된 질병 용어 모으기"""
    terms = set()
    for m in re.finditer(r"[‘'“\"「]([^’'”\"」]{1,20})[’'”\"」]\s*(?:이라|라)\s*함은", whole):
        t = clean(m.group(1))
        if any(nospace(t).endswith(e) for e in DISEASE_ENDINGS) and not re.search(r"\d|기간|동안|나이", t):
            terms.add(t)
    return terms


def known_terms(whole, codes):
    """지급사유에서 찾을 질병 용어 = 약관 정의 용어 + 급부코드명의 질병 부분"""
    terms = {nospace(t) for t in defined_terms(whole)}
    if codes is not None:
        terms |= {d for d in codes["_d"] if len(d) >= 2 and any(d.endswith(e) for e in DISEASE_ENDINGS)}
    terms |= set(TERM_ALIAS)
    return sorted(terms, key=len, reverse=True)


def terms_in_reason(reason, terms):
    """지급사유에 나온 질병 용어(긴 것 우선, 겹치지 않게). '(…제외)' 안과 'OO보장개시일' 은 제외"""
    t = re.sub(r"\([^()]*제외[^()]*\)", " ", reason)
    t = re.sub(r"\S*보장\s*개시일", " ", t)
    quoted = {nospace(q) for q in re.findall(r"[‘'“\"「]([^’'”\"」]{1,20})[’'”\"」]", t)
              if any(nospace(q).endswith(e) for e in DISEASE_ENDINGS)}
    terms = sorted(set(terms) | quoted, key=len, reverse=True)
    used, found = [False] * len(t), []
    for term in terms:
        pat = r"\s*".join(map(re.escape, term))
        for m in re.finditer(pat, t):
            s0, e0 = m.span()
            if any(used[s0:e0]):
                continue
            if len(term) <= 2 and s0 > 0 and re.match(r"[가-힣]", t[s0 - 1]):   # '암'은 단어 시작일 때만
                continue
            for k in range(s0, e0):
                used[k] = True
            found.append((s0, term))
    found.sort()
    out = []
    for _, term in found:
        if term not in out:
            out.append(term)
    return out


def _score(term, acts, row):
    """질병 용어 + 행위어 ↔ 급부코드명 점수 (0~1)"""
    t = TERM_ALIAS.get(term, term)
    d = row["_d"]
    if not d:
        s = 0.0
    elif t == d:
        s = 0.6
    elif d in t or t in d:
        s = 0.35 + 0.2 * min(len(d), len(t)) / max(len(d), len(t))
    else:
        s = 0.45 * difflib.SequenceMatcher(None, t, d).ratio()
    ca = row["_a"]
    if acts:
        hit = len([a for a in acts if a in ca]) / len(acts)
        s += 0.4 * hit if hit else (-0.1 if ca else 0.1)
    else:
        s += 0.2 if not ca else 0.1
    return s


def names_without_codes(name, reason, terms):
    """급부체계(급부코드표)가 없을 때: 약관 기준 급부명 만들기 → [("", 약관급부명, None, 근거)]
    - 지급사유에 질병명이 2개 이상 → 질병명마다 '질병명 + 행위' (예: 기타피부암 진단, 제자리암 진단)
    - 질병명이 1개이고 급부명에 들어 있거나, 질병명이 없으면 → 급부명 그대로 (예: 고액암진단)
    - 질병명이 1개인데 급부명에 없으면 → '질병명 + 행위' (예: 암진단비 → 일반암 진단)"""
    n = norm_benefit(name)
    acts = actions_in(name) or actions_in(reason)
    act = acts[0] if acts else ""
    found = terms_in_reason(reason, terms)

    def label(term):
        t = TERM_ALIAS.get(term, term)
        return f"{t} {act}".strip()

    if len(found) >= 2:
        return [("", label(t), None, f"지급사유 질병명: {t}") for t in found]
    if len(found) == 1 and (found[0] not in n or found[0] in TERM_ALIAS):
        return [("", label(found[0]), None, f"지급사유 질병명: {found[0]}")]
    return [("", display_name(name), None, "약관 급부명")]


def display_name(name):
    """표시용 급부명: 괄호·공백 제거, 끝의 '자금/비/금'은 앞이 행위어(진단·수술…)일 때만 뗌
    예) 고액암진단자금 → 고액암진단, 건강관리자금 → 건강관리자금"""
    s = nospace(re.sub(r"[\(\[【].*?[\)\]】]", "", name or ""))
    for suf in NAME_SUFFIXES:
        if s.endswith(suf) and any(s[: -len(suf)].endswith(a) for a in ACTION_WORDS):
            return s[: -len(suf)]
    return s


def match_codes(name, reason, codes, terms):
    """급부 1개 → [(급부코드, 급부코드명, 점수, 근거)] (여러 개일 수 있음)"""
    if codes is None or codes.empty:
        return names_without_codes(name, reason, terms)
    n = norm_benefit(name)
    # 1) 급부명 전체가 거의 같은 코드
    ratios = codes["_n"].map(lambda c: difflib.SequenceMatcher(None, n, c).ratio())
    i = ratios.idxmax()
    if ratios[i] >= NAME_DIRECT_SCORE:
        return [(codes.at[i, "급부코드"], codes.at[i, "급부코드명"], round(float(ratios[i]), 2), "급부명 일치")]
    # 2) 지급사유의 질병 용어마다
    acts = actions_in(name) or actions_in(reason)
    found = terms_in_reason(reason, terms)
    out = []
    for term in found:
        sc = codes.apply(lambda r: _score(term, acts, r), axis=1)
        j = sc.idxmax()
        if sc[j] >= MATCH_MIN_SCORE:
            out.append((codes.at[j, "급부코드"], codes.at[j, "급부코드명"], round(float(sc[j]), 2),
                        f"지급사유 질병명: {term}"))
        else:
            out.append(("", "", round(float(sc[j]), 2), f"지급사유 질병명: {term} (맞는 코드 없음)"))
    if any(c for c, *_ in out):
        return out
    # 3) 급부명과 가장 비슷한 코드
    d = disease_part(n)
    sc = codes.apply(lambda r: max(_score(d, acts, r), difflib.SequenceMatcher(None, n, r["_n"]).ratio()), axis=1)
    j = sc.idxmax()
    if sc[j] >= MATCH_MIN_SCORE:
        return [(codes.at[j, "급부코드"], codes.at[j, "급부코드명"], round(float(sc[j]), 2), "급부명 유사")]
    return out or [("", "", round(float(sc[j]), 2), "맞는 코드 없음")]


# =====================================================================
# 한 파일 처리
# =====================================================================


def product_title(text):
    """특약 제목이면 제목(번호 뗀 것)을, 아니면 None"""
    t = clean(text)
    if re.search(r"[·…\.ㆍ]{3,}", t):          # 목차 줄(점선) 제외
        return None
    if not t or len(t) > PRODUCT_TITLE_MAXLEN or re.match(r"^제\s*\d+\s*조", t):
        return None
    if any(w in t for w in PRODUCT_TITLE_EXCLUDE):
        return None
    core = re.sub(r"\s*약관$", "", t)
    stem = re.sub(r"\s*(?:Ⅰ|Ⅱ|Ⅲ|Ⅳ|Ⅴ|II|III|IV|V)$", "", core)     # 개정 특약의 끝 로마숫자 (예: OO특약Ⅱ)
    if not (any(stem.endswith(e) for e in PRODUCT_TITLE_ENDINGS)
            or any(re.search(p, t) for p in PRODUCT_TITLE_PATTERNS)):
        return None
    core = re.sub(r"^\s*(?:[\d]+(?:[\-\.]\d+)*\.?|[①-⑳]|[가-하]\.)\s*", "", core)   # 앞 번호 제거
    return core if len(nospace(core)) >= 3 else None


def _starts_new_table(blocks, a, b, product, title=""):
    """a~b 사이에 새 표 제목(별표/지급기준표)이나 다른 특약 제목이 있으면 True"""
    for k, d in blocks[a + 1:b]:
        if k != "p":
            continue
        if ("별표" in d or "지급기준표" in nospace(d)) and nospace(d) != nospace(title):
            return True
        pt = product_title(d)
        if pt and nospace(pt) != nospace(product):
            return True
    return False


def _only_page_noise(blocks, a, b):
    """a~b 사이가 쪽번호·머리말 같은 짧은 문단뿐이면 True (표가 다음 쪽으로 이어진 경우 판단용)"""
    between = blocks[a + 1:b]
    return len(between) <= 4 and all(k == "p" and len(d) <= 40 for k, d in between)


def extract_file(path, codes, log, word_reader):
    blocks, how = read_any(path, word_reader, log)
    whole = "\n".join(b[1] for b in blocks if b[0] == "p")
    terms = known_terms(whole, codes)          # 지급사유에서 찾을 질병 용어 (약관 정의 + 급부코드명)

    # 1) 지급기준표 모으기 (다음 쪽으로 이어진 표는 하나로 합침)
    file_product = os.path.splitext(os.path.basename(path))[0]
    groups, cur = [], None
    product, product_idx = "", 0          # 지금까지 나온 마지막 특약 제목
    for bi, (kind, data) in enumerate(blocks):
        if kind == "p":
            pt = product_title(data)
            if pt:
                product, product_idx = pt, bi
            continue
        if not data:
            continue
        hdr_i = next((ri for ri, row in enumerate(data[:3]) if _is_benefit_header(row)), None)
        width = len(data[0])
        is_cont = (cur is not None and _only_page_noise(blocks, cur["end"], bi)
                   and width == cur["width"]
                   and not _starts_new_table(blocks, cur["end"], bi, cur["product"], cur["title"]))
        if hdr_i is not None:
            body = data[hdr_i + 1:]
            if not is_cont:            # 새 지급기준표
                header = data[hdr_i]
                prev_p = [blocks[k][1] for k in range(bi - 1, max(-1, bi - 6), -1) if blocks[k][0] == "p"]
                title = next((p for p in prev_p if "별표" in p or "지급기준표" in nospace(p)),
                             prev_p[0] if prev_p else "")
                near = []                  # 표 바로 위 문단 (앞 표를 넘어가지 않음: 앞 표의 [기준: …] 을 가져오지 않도록)
                for k in range(bi - 1, max(-1, bi - 6), -1):
                    if blocks[k][0] != "p":
                        break
                    near.append(blocks[k][1])
                near += [" ".join(r) for r in data[:hdr_i]]
                base, base_src = find_base_amount(near)
                cur = {"title": title, "product": product or file_product,
                       "base": base, "base_src": base_src, "near": near,
                       "section": (product_idx, bi),
                       "width": width, "start": bi, "end": bi, "rows": [],
                       "c": (find_col(header, HEADER_BENEFIT), find_col(header, HEADER_REASON),
                             find_col(header, HEADER_AMOUNT)),
                       "hname": nospace(header[find_col(header, HEADER_BENEFIT)])}
                groups.append(cur)
        elif is_cont:                  # 머리글 없이 다음 쪽으로 이어진 부분
            body = data
        else:
            cur = None
            continue

        c_name, c_reason, c_amt = cur["c"]
        for j, row in enumerate(body):
            if max(c_name, c_reason, c_amt) >= len(row):
                continue
            name, reason, amount = row[c_name], row[c_reason], row[c_amt]
            if nospace(name) == cur["hname"]:
                continue
            if not nospace(name):
                if not cur["rows"] or not (nospace(reason) or nospace(amount)):
                    continue
                if is_cont and j == 0:  # 쪽이 바뀌며 한 칸이 둘로 잘린 경우 → 앞 행에 이어붙임
                    last = cur["rows"][-1]
                    last[1] = clean(last[1] + " " + reason)
                    last[2] = clean(last[2] + " " + amount)
                    continue
                name = cur["rows"][-1][0]  # 급부명 칸이 세로 병합된 경우
            cur["rows"].append([name, reason, amount])
        cur["end"] = bi

    # 2) 급부별 값 뽑기
    benefits = []
    for g in groups:
        # 표 아래 (주): 다음 표·다음 별표·다른 특약 제목이 나오기 전까지
        notes = []
        for k in range(g["end"] + 1, len(blocks)):
            kind, d = blocks[k]
            if kind == "tbl" or len(notes) >= 60 or d.startswith("별표"):
                break
            pt = product_title(d)
            if pt and nospace(pt) != nospace(g["product"]):
                break
            notes.append(d)
        notes = "\n".join(notes)
        # 이 특약 본문 (특약 제목 ~ 표): 'OO보장개시일' 정의를 이 특약 안에서 먼저 찾기 위함
        s0, s1 = g["section"]
        section = "\n".join(d for k, d in blocks[s0:s1] if k == "p")
        for name, reason, amount in g["rows"]:
            base, base_src = g["base"], g["base_src"]
            if base is None and find_money(amount):     # 표 근처에 없으면 특약 본문에서 '기준' 문구 찾기
                base, base_src = find_base_amount(section.split("\n") + notes.split("\n"))
            benefits.append(parse_benefit(path, g["product"], g["title"], name, reason, amount,
                                          notes, section, whole, codes, base, base_src, g["near"], terms))

    tag = {"Word": " (Word로 읽음)", "PDF": " (PDF)"}.get(how, "")
    n_prod = len({g["product"] for g in groups})
    log(f"  {os.path.basename(path)}: 특약 {n_prod}개, 급부 {len(benefits)}개{tag}")
    return benefits


def parse_benefit(path, product, title, raw_name, reason, amount, notes, section, whole, codes,
                  base=None, base_src="", near=(), terms=None):
    name, article = split_name(raw_name)
    both = reason + " " + amount
    memo = []

    unit, unit_src = None, ""
    if "구좌" in nospace(both):              # 1구좌 보험료: 지급금액·지급사유 → 표 근처 → 특약 본문 → (주) 순서로 찾음
        unit, unit_src = find_unit_premium([amount, reason] + list(near) + section.split("\n") + notes.split("\n"))
    basis, basis_word = get_basis(amount, unit)
    if basis in UNIT_PREMIUM_CODES.values():
        base, base_src = unit, unit_src      # 금액으로 적힌 지급금액은 1구좌 보험료로 나눠 지급율 계산
    if basis == "확인":
        memo.append(f"산출기준 코드 확인({basis_word or '문구 없음'})")
    rate, rate_memo = get_rate(amount, base)
    if rate_memo:
        memo.append(rate_memo)
    red_years, red_ratio, red_mode, red_src = get_reduction(both)
    wait, wait_src, wait_memo = get_wait_days(reason, notes, whole, section)
    if wait_memo and "문구 없음" not in wait_memo:
        memo.append(wait_memo)
    if wait is None and "문구 없음" in wait_memo:
        wait = 0
        wait_src = "보장개시일 문구 없음 → 0"
    pay_term, pay_term_src = get_pay_term(both)
    once, once_src = get_once_flag(name, both)
    if once is None:
        memo.append(f"적립금 지급여부 확인({once_src})")
    matches = match_codes(name, reason, codes, terms or [])
    for c, cn, sc, why in matches:
        if codes is None:
            break
        if not c:
            memo.append(f"급부코드 못 찾음({why})")
        elif sc < MATCH_OK_SCORE:
            memo.append(f"급부코드 매칭 확인({cn}, {sc})")
    if red_years and "갱신" in notes and "감액" in notes:
        memo.append("감액: 갱신계약 미적용 문구 있음(갱신형 확인)")

    return {
        "파일명": os.path.basename(path), "특약명": product, "별표명": title,
        "급부명": name, "조항": article,
        "급부코드": ", ".join(c for c, *_ in matches if c),
        "급부코드명": ", ".join(cn for c, cn, *_ in matches if c),
        "약관급부명": ", ".join(cn for c, cn, *_ in matches if not c and cn),
        "매칭점수": ", ".join(str(sc) for c, cn, sc, _ in matches if c),
        "매칭근거": " / ".join(w for *_, w in matches),
        "_matches": matches,
        "납입주기": get_pay_cycle(amount),
        "산출기준": basis, "기본지급율": rate, "지급주기": pay_term, "기준금액(원)": base, "근거_기준금액": base_src,
        "감액기간(년)": red_years, "감액비율": red_ratio, "감액방식": red_mode,
        "면책기간": wait, "적립금지급여부": once,
        "확인필요": " / ".join(memo),
        "근거_산출기준": basis_word, "근거_지급주기": pay_term_src, "근거_감액": red_src, "근거_면책": wait_src, "근거_적립금": once_src,
        "원문_지급사유": reason, "원문_지급금액": amount,
    }


def expand_rows(b):
    """급부 1개 → 경과 구간별 행"""
    base = {k: b[k] for k in ["파일명", "특약명", "급부명", "조항", "급부코드", "급부코드명",
                              "산출기준", "지급주기", "적립금지급여부", "면책기간"]}
    rate = b["기본지급율"]
    if rate is not None and b["감액기간(년)"]:
        cut = 12 * b["감액기간(년)"]
        reduced = b["감액비율"] if b["감액방식"] == "절대값" else rate * b["감액비율"]
        segs = [(0, cut, round(reduced, 4)), (cut + 1, END_MONTH, rate)]
    else:
        segs = [(0, END_MONTH, rate)]
    out = []
    matches = [m for m in b.get("_matches") or [] if m[0] or m[1]] or [("", "", None, "")]
    for code, code_name, score, why in matches:      # 급부코드(또는 약관급부명)가 여러 개면 그만큼 행을 만듦
        for i, (s, e, r) in enumerate(segs):
            row = dict(base)
            row.update({"급부코드": code, "급부코드명": code_name if code else "",
                        "약관급부명": "" if code else code_name, "매칭점수": score, "매칭근거": why,
                        "경과년1": s, "경과년2": e, "지급율": r, "확인필요": b["확인필요"]})
            if i > 0:
                row["면책기간"] = 0      # 면책은 첫 경과 구간(예: 0~12)에만. 이후 구간은 이미 면책기간이 지남
            out.append(row)
    return out


# =====================================================================
# 전체 실행
# =====================================================================


def list_word_files(src):
    if os.path.isfile(src):
        return [src]
    files = []
    for dp, _, fns in os.walk(src):
        for fn in fns:
            if fn.startswith("~$"):
                continue
            if fn.lower().endswith((".docx", ".docm", ".doc", ".pdf")):
                files.append(os.path.join(dp, fn))
    return sorted(files)


def run(src, out_dir, code_xlsx=None, log=print, stop_event=None):
    codes = load_code_table(code_xlsx) if code_xlsx else None
    files = list_word_files(src)
    if not files:
        raise RuntimeError("Word/PDF 파일을 찾지 못했습니다.")
    log(f"대상 파일 {len(files)}개")

    benefits, status = [], []
    stopped = False
    reader = WordReader()
    try:
        for i, f in enumerate(files, 1):
            if stop_event is not None and stop_event.is_set():
                stopped = True
                log(f"중단: {i - 1}/{len(files)}개까지 처리한 결과를 저장합니다.")
                for rest in files[i - 1:]:
                    status.append({"파일명": os.path.basename(rest), "결과": "중단(미처리)",
                                   "급부수": 0, "경로": rest})
                break
            log(f"[{i}/{len(files)}]")
            try:
                got = extract_file(f, codes, log, reader)
                benefits += got
                status.append({"파일명": os.path.basename(f), "결과": "성공" if got else "지급기준표 못 찾음",
                               "급부수": len(got), "경로": f})
            except Exception as e:
                log(f"  실패: {os.path.basename(f)} -> {e}")
                status.append({"파일명": os.path.basename(f), "결과": f"실패: {e}", "급부수": 0, "경로": f})
    finally:
        reader.quit()       # 뒤에서 연 Word 닫기

    draft_cols = ["파일명", "특약명", "급부명", "조항", "급부코드", "급부코드명", "매칭점수", "매칭근거",
                  "산출기준", "지급율", "지급주기", "적립금지급여부", "면책기간", "경과년1", "경과년2", "확인필요"]
    draft = pd.DataFrame([r for b in benefits for r in expand_rows(b)], columns=draft_cols)
    if codes is None:
        draft.insert(draft.columns.get_loc("급부코드"), "약관급부명",
                     [r.get("약관급부명", "") for b in benefits for r in expand_rows(b)])
        draft = draft.drop(columns=["급부코드", "급부코드명", "매칭점수"])
    for c in ["적립금지급여부", "면책기간", "경과년1", "경과년2"]:   # 1.0 → 1 로 표시
        draft[c] = pd.to_numeric(draft[c], errors="coerce").astype("Int64")
    raw = pd.DataFrame(benefits).drop(columns=["_matches"], errors="ignore")
    if codes is None:
        raw = raw.drop(columns=["급부코드", "급부코드명", "매칭점수"], errors="ignore")
    else:
        raw = raw.drop(columns=["약관급부명"], errors="ignore")
    stat = pd.DataFrame(status)
    if len(raw):
        prod = (raw.assign(_chk=raw["확인필요"] != "")
                .groupby(["파일명", "특약명"], sort=False)
                .agg(별표명=("별표명", "first"), 급부수=("급부명", "size"), 확인필요급부=("_chk", "sum"),
                     급부목록=("급부명", lambda s: ", ".join(dict.fromkeys(s))))
                .reset_index())
    else:
        prod = pd.DataFrame(columns=["파일명", "특약명", "별표명", "급부수", "확인필요급부", "급부목록"])

    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S") + ("_중단" if stopped else "")
    out = os.path.join(out_dir, f"통계테이블초안_{stamp}.xlsx")
    n = 2
    while os.path.exists(out):
        out = os.path.join(out_dir, f"통계테이블초안_{stamp}_{n}.xlsx")
        n += 1
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        draft.to_excel(xw, sheet_name="초안", index=False)
        prod.to_excel(xw, sheet_name="특약별요약", index=False)
        raw.to_excel(xw, sheet_name="급부원문", index=False)
        stat.to_excel(xw, sheet_name="처리결과", index=False)
        from openpyxl.styles import PatternFill, Alignment
        yellow = PatternFill("solid", fgColor="FFF2CC")
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(8, w * 1.5), 60)
        ws = xw.book["초안"]
        chk = [c.value for c in ws[1]].index("확인필요") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, chk).value:
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).fill = yellow
        for name in ["급부원문"]:
            for row in xw.book[name].iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")

    n_chk = int((draft["확인필요"] != "").sum()) if len(draft) else 0
    log(f"급부 {len(benefits)}개 → 초안 {len(draft)}행 (확인필요 {n_chk}행)")
    log(f"결과 파일: {out}")
    return out


# =====================================================================
# 화면
# =====================================================================

def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("약관 지급기준표 추출 (통계테이블 초안)")
    root.geometry("780x500")
    root.lift()
    root.attributes("-topmost", True)
    root.after(800, lambda: root.attributes("-topmost", False))

    v_src, v_code, v_out = tk.StringVar(), tk.StringVar(), tk.StringVar()
    q = queue.Queue()

    def pick_dir(var):
        d = filedialog.askdirectory()
        if d:
            var.set(d)
            if var is v_src and not v_out.get():
                v_out.set(d)

    def pick_file(var, types):
        f = filedialog.askopenfilename(filetypes=types)
        if f:
            var.set(f)
            if var is v_src and not v_out.get():
                v_out.set(os.path.dirname(f))

    word_t = [("Word/PDF", "*.docx *.doc *.docm *.pdf")]
    xl_t = [("Excel/CSV", "*.xlsx *.xls *.csv")]
    items = [
        ("약관 폴더/파일", v_src, [("폴더", lambda: pick_dir(v_src)), ("파일", lambda: pick_file(v_src, word_t))]),
        ("급부코드표(선택)", v_code, [("파일", lambda: pick_file(v_code, xl_t))]),
        ("결과 저장 위치", v_out, [("폴더", lambda: pick_dir(v_out))]),
    ]
    for r, (label, var, btns) in enumerate(items):
        tk.Label(root, text=label, width=15, anchor="w").grid(row=r, column=0, padx=8, pady=6, sticky="w")
        tk.Entry(root, textvariable=var, width=62).grid(row=r, column=1, padx=4, sticky="ew")
        f = tk.Frame(root)
        f.grid(row=r, column=2, padx=6)
        for t, cmd in btns:
            tk.Button(f, text=t, command=cmd, width=5).pack(side="left", padx=2)

    logbox = scrolledtext.ScrolledText(root, height=18)
    logbox.grid(row=4, column=0, columnspan=3, padx=8, pady=8, sticky="nsew")
    root.grid_rowconfigure(4, weight=1)
    root.grid_columnconfigure(1, weight=1)

    def log(m):
        q.put(("log", m))

    def poll():
        try:
            while True:
                kind, m = q.get_nowait()
                if kind == "log":
                    logbox.insert("end", m + "\n")
                    logbox.see("end")
                elif kind == "done":
                    btn.config(state="normal")
                    stop_btn.config(state="disabled")
                    msg = ("중단했습니다. 지금까지 처리한 결과를 저장했어요." if stop_event.is_set()
                           else "초안을 만들었습니다.")
                    if messagebox.askyesno("완료", msg + " 엑셀을 열까요?"):
                        try:
                            os.startfile(m)
                        except Exception:
                            pass
                elif kind == "error":
                    btn.config(state="normal")
                    stop_btn.config(state="disabled")
                    messagebox.showerror("오류", m)
        except queue.Empty:
            pass
        root.after(200, poll)

    def start():
        src, code, out = v_src.get().strip(), v_code.get().strip(), v_out.get().strip()
        if not src or not os.path.exists(src):
            messagebox.showwarning("확인", "약관 폴더 또는 파일을 선택해 주세요.")
            return
        if code and not os.path.isfile(code):
            messagebox.showwarning("확인", "급부코드표 파일을 확인해 주세요.")
            return
        if not out or not os.path.isdir(out):
            messagebox.showwarning("확인", "결과 저장 위치를 선택해 주세요.")
            return
        btn.config(state="disabled")
        stop_btn.config(state="normal")
        stop_event.clear()
        logbox.delete("1.0", "end")
        log("추출 시작...")

        def work():
            try:
                q.put(("done", run(src, out, code or None, log, stop_event)))
            except Exception as e:
                q.put(("error", str(e)))

        threading.Thread(target=work, daemon=True).start()

    def stop():
        stop_event.set()
        stop_btn.config(state="disabled")
        log("중단 요청 → 지금 읽는 파일까지 마치고 멈춥니다...")

    stop_event = threading.Event()
    bar = tk.Frame(root)
    bar.grid(row=3, column=0, columnspan=3, pady=6)
    btn = tk.Button(bar, text="실행", width=16, command=start)
    btn.pack(side="left", padx=6)
    stop_btn = tk.Button(bar, text="중단 (지금까지 저장)", width=20, command=stop, state="disabled")
    stop_btn.pack(side="left", padx=6)
    poll()
    root.mainloop()


if __name__ == "__main__":
    run_gui()
