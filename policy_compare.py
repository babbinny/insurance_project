# -*- coding: utf-8 -*-
"""
약관 변경 전·후 비교 프로그램 (지급기준표 중심)
----------------------------------------------------------------
- 같은 상품의 '변경 전' 약관과 '변경 후' 약관(Word/PDF)을 넣으면 달라진 점을 정리합니다.
- 지급기준표는 benefit_extractor2.py 와 같은 규칙으로 읽어서 급부마다 비교합니다.
    지급사유 · 지급금액 · 산출기준 · 지급율 · 지급주기 · 감액 · 면책기간 · 적립금 지급여부
    → 급부마다 '변경 / 신설 / 삭제 / 동일'
- 특약(상품) 목록이 바뀐 것(신설·삭제)과, 지급기준표 밖 본문 문장의 변경(조항별)도 함께 정리합니다.
- 결과: 약관비교_날짜시간.xlsx
    [요약]           건수 요약
    [지급기준표_비교] 급부별 변경 내용 (변경 노란색, 신설 파란색, 삭제 빨간색)
    [특약_비교]      특약(상품) 목록 비교
    [본문_변경]      지급기준표 밖 본문에서 바뀐 문장 (조항별)
- benefit_extractor2.py 와 같은 폴더에 두고, Spyder에서 F5로 실행하면 창이 뜹니다.
"""
import os
import re
import queue
import difflib
import datetime
import threading
from collections import OrderedDict

import pandas as pd

import benefit_extractor2 as be

# =====================================================================
# 설정
# =====================================================================

# 비교할 항목: (표시 이름, 급부 정보의 키)
COMPARE_FIELDS = [
    ("지급사유", "원문_지급사유"),
    ("지급금액", "원문_지급금액"),
    ("산출기준", "산출기준"),
    ("지급율", "기본지급율"),
    ("지급주기", "지급주기"),
    ("감액", "_감액"),
    ("면책기간", "면책기간"),
    ("적립금지급여부", "적립금지급여부"),
]
PRODUCT_MATCH_SCORE = 0.6     # 특약명이 이 점수 이상 비슷하면 같은 특약으로 봄 (예: 'Ⅱ' 붙은 개정 특약)
BENEFIT_MATCH_SCORE = 0.75    # 급부명이 이 점수 이상 비슷하면 같은 급부로 봄
BODY_MIN_CHARS = 15           # 본문 비교에서 이보다 짧은 줄(쪽번호·머리말 등)은 뺌
BODY_PAIR_SCORE = 0.5         # 바뀐 문장끼리 이 점수 이상 비슷하면 '변경'(아니면 삭제+추가)
BODY_MAX_ROWS = 2000          # 본문 변경은 최대 이만큼만 적음

COLORS = {"변경": "FFF2CC", "신설": "DDEBF7", "삭제": "F8CBAD", "추가": "DDEBF7"}

# =====================================================================
# 도우미
# =====================================================================


def norm_product(s):
    """특약명 비교용: 공백·'무배당/(무)'·'약관'·끝의 로마숫자 제거"""
    s = be.nospace(s or "")
    s = re.sub(r"무배당|\(무\)|약관$", "", s)
    s = re.sub(r"(?:Ⅰ|Ⅱ|Ⅲ|Ⅳ|Ⅴ|II|III|IV)+$", "", s)
    return s


def ratio(a, b):
    return difflib.SequenceMatcher(None, a, b).ratio()


def fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def word_diff(a, b):
    """두 문장의 단어 단위 차이 → [(구분, 글자)] 구분: 'eq' / 'del' / 'ins'"""
    ta, tb = (a or "").split(), (b or "").split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ta, tb, autojunk=False).get_opcodes():
        if op == "equal":
            out.append(("eq", " ".join(ta[i1:i2])))
        else:
            if i2 > i1:
                out.append(("del", " ".join(ta[i1:i2])))
            if j2 > j1:
                out.append(("ins", " ".join(tb[j1:j2])))
    return out


def diff_text(a, b):
    """엑셀용 한 줄 요약: 바뀐 단어만 '[-전-] [+후+]' 로"""
    parts = [f"[-{t}-]" if k == "del" else f"[+{t}+]" for k, t in word_diff(a, b) if k != "eq"]
    return " ".join(parts)


def _reduction(b):
    if not b.get("감액기간(년)"):
        return ""
    return f"{fmt(b['감액기간(년)'])}년 미만 {fmt(b['감액비율'])}({b.get('감액방식', '')})"


# =====================================================================
# 짝 맞추기
# =====================================================================


def pair_by_name(left, right, key, min_score):
    """이름이 같은 것끼리 순서대로 짝짓고, 남은 것은 가장 비슷한 것끼리(min_score 이상) 짝지음
    → [(왼쪽 or None, 오른쪽 or None)]"""
    pairs, used_r = [], set()
    by_key = OrderedDict()
    for j, r in enumerate(right):
        by_key.setdefault(key(r), []).append(j)
    left_rest = []
    for i, l in enumerate(left):
        cand = [j for j in by_key.get(key(l), []) if j not in used_r]
        if cand:
            used_r.add(cand[0])
            pairs.append((i, cand[0]))
        else:
            left_rest.append(i)
    scored = sorted(((ratio(key(left[i]), key(right[j])), i, j)
                     for i in left_rest for j in range(len(right)) if j not in used_r), reverse=True)
    done_l = set()
    for sc, i, j in scored:
        if sc < min_score:
            break
        if i in done_l or j in used_r:
            continue
        done_l.add(i)
        used_r.add(j)
        pairs.append((i, j))
    matched_l = {i for i, _ in pairs}
    out = [(left[i], right[j]) for i, j in sorted(pairs)]
    out += [(left[i], None) for i in range(len(left)) if i not in matched_l]
    out += [(None, right[j]) for j in range(len(right)) if j not in used_r]
    return out


def group_products(benefits):
    groups = OrderedDict()
    for b in benefits:
        groups.setdefault(b["특약명"], []).append(b)
    return [{"name": k, "benefits": v} for k, v in groups.items()]


# =====================================================================
# 본문 비교 (지급기준표 밖 문장)
# =====================================================================


def body_lines(blocks):
    """문단 → [(조항, 특약, 문장)] (짧은 줄·쪽번호 제외)"""
    out, article, product = [], "", ""
    for kind, d in blocks:
        if kind != "p":
            continue
        pt = be.product_title(d)
        if pt:
            product, article = pt, ""
        m = re.match(r"^(제\s*\d+\s*조(?:의\s*\d+)?(?:\s*\([^)]{1,30}\))?)", d)
        if m:
            article = be.clean(m.group(1))
        if len(be.nospace(d)) < BODY_MIN_CHARS or re.fullmatch(r"[\d\s\-–/.()]+", d):
            continue
        out.append((article, product, d))
    return out


def compare_body(blocks_a, blocks_b):
    la, lb = body_lines(blocks_a), body_lines(blocks_b)
    na, nb = [be.nospace(x[2]) for x in la], [be.nospace(x[2]) for x in lb]
    rows = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, na, nb, autojunk=False).get_opcodes():
        if op == "equal":
            continue
        a_idx, b_idx = list(range(i1, i2)), list(range(j1, j2))
        # 바뀐 구간 안에서 비슷한 문장끼리 '변경'으로 짝지음
        while a_idx and b_idx:
            i, j = a_idx[0], b_idx[0]
            if ratio(na[i], nb[j]) >= BODY_PAIR_SCORE:
                rows.append(("변경", lb[j][1] or la[i][1], lb[j][0] or la[i][0], la[i][2], lb[j][2]))
                a_idx.pop(0)
                b_idx.pop(0)
            elif len(a_idx) >= len(b_idx):
                rows.append(("삭제", la[i][1], la[i][0], la[i][2], ""))
                a_idx.pop(0)
            else:
                rows.append(("추가", lb[j][1], lb[j][0], "", lb[j][2]))
                b_idx.pop(0)
        rows += [("삭제", la[i][1], la[i][0], la[i][2], "") for i in a_idx]
        rows += [("추가", lb[j][1], lb[j][0], "", lb[j][2]) for j in b_idx]
    return rows


# =====================================================================
# 전체 비교
# =====================================================================


def read_both(path_a, path_b, log):
    """두 약관을 benefit_extractor2 로 읽음 → (급부들 A, 급부들 B, 문단 A, 문단 B)"""
    cache = {}
    orig = be.read_any

    def cached(path, word_reader, lg=None):          # 지급기준표 추출과 본문 비교가 같은 읽기 결과를 씀
        if path not in cache:
            cache[path] = orig(path, word_reader, lg)
        return cache[path]

    be.read_any = cached
    reader = be.WordReader()
    try:
        log("[변경 전]")
        ben_a = be.extract_file(path_a, None, log, reader)
        log("[변경 후]")
        ben_b = be.extract_file(path_b, None, log, reader)
    finally:
        be.read_any = orig
        reader.quit()
    return ben_a, ben_b, cache[path_a][0], cache[path_b][0]


def compare_benefits(ben_a, ben_b):
    """→ (특약 비교 행들, 급부 비교 항목들)"""
    for b in ben_a + ben_b:
        b["_감액"] = _reduction(b)
    prods = pair_by_name(group_products(ben_a), group_products(ben_b),
                         lambda p: norm_product(p["name"]), PRODUCT_MATCH_SCORE)
    prod_rows, items = [], []
    for pa, pb in prods:
        if pa and pb:
            pstatus = "동일" if pa["name"] == pb["name"] else "이름 변경"
        else:
            pstatus = "신설" if pb else "삭제"
        bens = pair_by_name(pa["benefits"] if pa else [], pb["benefits"] if pb else [],
                            lambda b: be.norm_benefit(b["급부명"]), BENEFIT_MATCH_SCORE)
        n = {"변경": 0, "신설": 0, "삭제": 0, "동일": 0}
        for a, b in bens:
            changes = []
            if a and b:
                for label, key in COMPARE_FIELDS:
                    va, vb = fmt(a.get(key)), fmt(b.get(key))
                    if be.nospace(va) != be.nospace(vb):
                        changes.append({"field": label, "before": va, "after": vb, "diff": word_diff(va, vb)})
                status = "변경" if changes or a["급부명"] != b["급부명"] else "동일"
                if a["급부명"] != b["급부명"]:
                    changes.insert(0, {"field": "급부명", "before": a["급부명"], "after": b["급부명"],
                                       "diff": word_diff(a["급부명"], b["급부명"])})
            else:
                status = "신설" if b else "삭제"
            n[status] += 1
            x = a or b
            items.append({
                "status": status,
                "product_before": pa["name"] if pa else "", "product_after": pb["name"] if pb else "",
                "name": (b or a)["급부명"], "article_before": a["조항"] if a else "", "article_after": b["조항"] if b else "",
                "changes": changes,
                "before": {label: fmt(a.get(key)) for label, key in COMPARE_FIELDS} if a else None,
                "after": {label: fmt(b.get(key)) for label, key in COMPARE_FIELDS} if b else None,
                "table_before": a["별표명"] if a else "", "table_after": b["별표명"] if b else "",
                "_x": x,
            })
        prod_rows.append({"특약명(변경 전)": pa["name"] if pa else "", "특약명(변경 후)": pb["name"] if pb else "",
                          "상태": pstatus, "급부 변경": n["변경"], "급부 신설": n["신설"],
                          "급부 삭제": n["삭제"], "급부 동일": n["동일"]})
    return prod_rows, items


def compare(path_a, path_b, out_dir=None, log=print):
    """두 약관 비교. out_dir 를 주면 엑셀 저장 → (엑셀 경로 or None, 결과 dict)"""
    for p in (path_a, path_b):
        if not os.path.isfile(p):
            raise RuntimeError(f"파일을 찾지 못했습니다: {p}")
    ben_a, ben_b, blocks_a, blocks_b = read_both(path_a, path_b, log)
    if not ben_a and not ben_b:
        log("  (참고) 두 약관 모두에서 지급기준표를 찾지 못했습니다. 본문만 비교합니다.")
    prod_rows, items = compare_benefits(ben_a, ben_b)
    body = compare_body(blocks_a, blocks_b)
    counts = OrderedDict((k, sum(1 for x in items if x["status"] == k)) for k in ("변경", "신설", "삭제", "동일"))
    pcounts = OrderedDict((k, sum(1 for r in prod_rows if r["상태"] == k)) for k in ("신설", "삭제", "이름 변경"))
    bcounts = OrderedDict((k, sum(1 for r in body if r[0] == k)) for k in ("변경", "추가", "삭제"))
    log(f"급부: 변경 {counts['변경']} / 신설 {counts['신설']} / 삭제 {counts['삭제']} / 동일 {counts['동일']}")
    log(f"특약: 신설 {pcounts['신설']} / 삭제 {pcounts['삭제']} / 이름 변경 {pcounts['이름 변경']}")
    log(f"본문 문장: 변경 {bcounts['변경']} / 추가 {bcounts['추가']} / 삭제 {bcounts['삭제']}")

    result = {
        "file_before": os.path.basename(path_a), "file_after": os.path.basename(path_b),
        "counts": counts, "product_counts": pcounts, "body_counts": bcounts,
        "products": prod_rows,
        "benefits": [{k: v for k, v in x.items() if k != "_x"} for x in items],
        "body": [{"kind": k, "product": p, "article": a, "before": bf, "after": af,
                  "diff": word_diff(bf, af) if k == "변경" else None} for k, p, a, bf, af in body[:BODY_MAX_ROWS]],
        "body_total": len(body),
    }
    out = save_excel(result, out_dir) if out_dir else None
    if out:
        log(f"결과 파일: {out}")
    return out, result


# =====================================================================
# 엑셀 저장
# =====================================================================


def save_excel(result, out_dir):
    from openpyxl.styles import PatternFill, Alignment, Font
    summary = [["변경 전 파일", result["file_before"]], ["변경 후 파일", result["file_after"]], []]
    summary += [["급부 " + k, v] for k, v in result["counts"].items()]
    summary += [["특약 " + k, v] for k, v in result["product_counts"].items()]
    summary += [["본문 문장 " + k, v] for k, v in result["body_counts"].items()]

    ben_rows = []
    order = {"변경": 0, "신설": 1, "삭제": 2, "동일": 3}
    for x in sorted(result["benefits"], key=lambda x: order[x["status"]]):
        row = OrderedDict([("상태", x["status"]), ("특약명", x["product_after"] or x["product_before"]),
                           ("급부명", x["name"]), ("조항(전)", x["article_before"]), ("조항(후)", x["article_after"]),
                           ("변경 항목", ", ".join(c["field"] for c in x["changes"])),
                           ("변경 내용", " / ".join(f"{c['field']}: {diff_text(c['before'], c['after'])}"
                                                   for c in x["changes"]))])
        for label, _ in COMPARE_FIELDS:
            row[f"{label}(전)"] = (x["before"] or {}).get(label, "")
            row[f"{label}(후)"] = (x["after"] or {}).get(label, "")
        ben_rows.append(row)
    ben_cols = ["상태", "특약명", "급부명", "조항(전)", "조항(후)", "변경 항목", "변경 내용"] + \
               [f"{l}({w})" for l, _ in COMPARE_FIELDS for w in ("전", "후")]

    body_rows = [OrderedDict([("구분", r["kind"]), ("특약", r["product"]), ("조항", r["article"]),
                              ("변경 전", r["before"]), ("변경 후", r["after"]),
                              ("바뀐 부분", diff_text(r["before"], r["after"]) if r["kind"] == "변경" else "")])
                 for r in result["body"]]
    body_cols = ["구분", "특약", "조항", "변경 전", "변경 후", "바뀐 부분"]

    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(out_dir, f"약관비교_{stamp}.xlsx")
    n = 2
    while os.path.exists(out):
        out = os.path.join(out_dir, f"약관비교_{stamp}_{n}.xlsx")
        n += 1
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        pd.DataFrame(summary, columns=["항목", "값"]).to_excel(xw, sheet_name="요약", index=False)
        pd.DataFrame(ben_rows, columns=ben_cols).to_excel(xw, sheet_name="지급기준표_비교", index=False)
        pd.DataFrame(result["products"], columns=["특약명(변경 전)", "특약명(변경 후)", "상태", "급부 변경", "급부 신설",
                                                  "급부 삭제", "급부 동일"]).to_excel(xw, sheet_name="특약_비교", index=False)
        pd.DataFrame(body_rows, columns=body_cols).to_excel(xw, sheet_name="본문_변경", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
            for col in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(8, w * 1.4), 60)
        for name, status_col in (("지급기준표_비교", 1), ("특약_비교", 3), ("본문_변경", 1)):
            ws = xw.book[name]
            for r in range(2, ws.max_row + 1):
                color = COLORS.get(ws.cell(r, status_col).value)
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(r, c)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if color:
                        cell.fill = PatternFill("solid", fgColor=color)
    return out


# =====================================================================
# 화면
# =====================================================================

def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("약관 변경 전·후 비교 (지급기준표 중심)")
    root.geometry("780x460")
    root.lift()
    root.attributes("-topmost", True)
    root.after(800, lambda: root.attributes("-topmost", False))
    v_a, v_b, v_out = tk.StringVar(), tk.StringVar(), tk.StringVar()
    q = queue.Queue()
    types = [("Word/PDF", "*.docx *.doc *.docm *.pdf")]

    def pick_file(var):
        f = filedialog.askopenfilename(filetypes=types)
        if f:
            var.set(f)
            if not v_out.get():
                v_out.set(os.path.dirname(f))

    def pick_dir():
        d = filedialog.askdirectory()
        if d:
            v_out.set(d)

    for r, (label, var, cmd) in enumerate([("변경 전 약관", v_a, lambda: pick_file(v_a)),
                                           ("변경 후 약관", v_b, lambda: pick_file(v_b)),
                                           ("결과 저장 위치", v_out, pick_dir)]):
        tk.Label(root, text=label, width=14, anchor="w").grid(row=r, column=0, padx=8, pady=6, sticky="w")
        tk.Entry(root, textvariable=var, width=66).grid(row=r, column=1, padx=4, sticky="ew")
        tk.Button(root, text="찾아보기", command=cmd).grid(row=r, column=2, padx=8)

    logbox = scrolledtext.ScrolledText(root, height=15)
    logbox.grid(row=4, column=0, columnspan=3, padx=8, pady=8, sticky="nsew")
    root.grid_rowconfigure(4, weight=1)
    root.grid_columnconfigure(1, weight=1)

    def log(m):
        q.put(("log", m))

    def poll():
        try:
            while True:
                kind, m = q.get_nowait()
                if kind == "log":
                    logbox.insert("end", m + "\n")
                    logbox.see("end")
                elif kind == "done":
                    btn.config(state="normal")
                    if messagebox.askyesno("완료", "비교가 끝났습니다. 결과 엑셀을 열까요?"):
                        try:
                            os.startfile(m)
                        except Exception:
                            pass
                elif kind == "error":
                    btn.config(state="normal")
                    messagebox.showerror("오류", m)
        except queue.Empty:
            pass
        root.after(200, poll)

    def start():
        a, b, out = v_a.get().strip(), v_b.get().strip(), v_out.get().strip()
        if not os.path.isfile(a) or not os.path.isfile(b):
            messagebox.showwarning("확인", "변경 전·후 약관 파일을 모두 선택해 주세요.")
            return
        if not os.path.isdir(out):
            messagebox.showwarning("확인", "결과 저장 위치를 선택해 주세요.")
            return
        btn.config(state="disabled")
        logbox.delete("1.0", "end")
        log("비교 시작...")

        def work():
            try:
                q.put(("done", compare(a, b, out, log)[0]))
            except Exception as e:
                q.put(("error", str(e)))

        threading.Thread(target=work, daemon=True).start()

    btn = tk.Button(root, text="비교 실행", width=16, command=start)
    btn.grid(row=3, column=0, columnspan=3, pady=6)
    poll()
    root.mainloop()


if __name__ == "__main__":
    run_gui()
