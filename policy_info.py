# -*- coding: utf-8 -*-
"""
약관 정보 요약 프로그램 (급부별 보험금 · 면책기간 · 적립금)
----------------------------------------------------------------
- 상품 약관(Word/PDF)을 넣으면 특약별로, 급부마다
    언제(지급사유) · 얼마(지급금액) · 면책기간 · 적립금 지급 여부 를 한눈에 정리합니다.
- 지급기준표는 benefit_extractor2.py 와 같은 규칙으로 읽습니다 (웹사이트 [약관 정보] 탭과 같은 결과).
- 결과: 약관정보_날짜시간.xlsx
    [약관정보]  급부별 요약 (확인필요 행은 노란색)
    [특약목록]  특약마다 급부 수 · 확인필요 수
    [처리결과]  파일별 성공/실패
- benefit_extractor2.py 와 같은 폴더에 두고, Spyder에서 F5로 실행하면 창이 뜹니다.
"""
import os
import queue
import datetime
import threading
from collections import OrderedDict

import pandas as pd

import benefit_extractor2 as be

# =====================================================================
# 설정
# =====================================================================

# 산출기준 코드 → 화면에 보일 이름 (benefit_extractor2.CALC_BASIS_CODES 와 같은 코드)
BASIS_LABELS = {
    "01": "보험가입금액", "02": "보험료", "03": "기납입보험료",
    "10": "MAX(가입금액, 기납입보험료) × 추가설계 비율",
    "95": "1구좌(보험료 1천만원)", "96": "1구좌(보험료 1억원)", "97": "1구좌(보험료 5천만원)", "98": "1구좌(보험료 100만원)",
    "확인": "산출기준 확인 필요",
}
UNIT_CODES = {"95", "96", "97", "98"}
CHECK_FILL = "FFF2CC"        # 확인필요 행 색

# =====================================================================
# 요약 만들기
# =====================================================================


def _pct(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "?"
    return f"{round(float(x) * 100, 2):g}%"


def _won(v):
    v = float(v)
    return f"{v / 10000:,.0f}만원" if v % 10000 == 0 else f"{v:,.0f}원"


def table_values(b):
    """통계테이블에 들어가는 값 → '산출기준 01 보험가입금액 × 100% · 가입 1년 미만 50% · 지급주기 100'"""
    code = b.get("산출기준") or "?"
    parts = [f"산출기준 {'' if code == '확인' else code + ' '}{BASIS_LABELS.get(code, '')} × {_pct(b.get('기본지급율'))}"]
    if b.get("기준금액(원)") and code not in UNIT_CODES:
        parts.append("기준 가입금액 " + _won(b["기준금액(원)"]))
    if b.get("감액기간(년)"):
        rate = b.get("기본지급율") or 0
        reduced = b["감액비율"] if b.get("감액방식") == "절대값" else rate * b["감액비율"]
        parts.append(f"가입 {b['감액기간(년)']}년 미만 {_pct(reduced)}")
    parts.append("지급주기 211 (n회 월지급)" if b.get("지급주기") == "211" else "지급주기 100")
    return " · ".join(parts)


def wait_text(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "확인필요"
    return f"{int(v)}일" if v > 0 else "없음"


def reserve_text(v):
    return {1: "지급", 0: "미지급"}.get(v, "확인필요")


def summarize(src, log=print):
    """약관 파일 또는 폴더 → (급부별 요약 행들, 특약 목록 행들, 처리결과 행들)"""
    files = be.list_word_files(src)
    if not files:
        raise RuntimeError("Word/PDF 파일을 찾지 못했습니다.")
    log(f"대상 파일 {len(files)}개")
    rows, status = [], []
    reader = be.WordReader()
    try:
        for i, f in enumerate(files, 1):
            log(f"[{i}/{len(files)}]")
            try:
                bens = be.extract_file(f, None, log, reader)
            except Exception as e:
                log(f"  실패: {os.path.basename(f)} -> {e}")
                status.append({"파일명": os.path.basename(f), "결과": f"실패: {e}", "급부수": 0})
                continue
            status.append({"파일명": os.path.basename(f), "결과": "성공" if bens else "지급기준표 못 찾음",
                           "급부수": len(bens)})
            for b in bens:
                rows.append(OrderedDict([
                    ("파일명", b["파일명"]), ("특약명", b["특약명"]), ("급부명", b["급부명"]), ("조항", b["조항"]),
                    ("지급사유 (언제)", b["원문_지급사유"]),
                    ("지급금액 (얼마)", b["원문_지급금액"]),
                    ("통계테이블 값", table_values(b)),
                    ("면책기간", wait_text(b["면책기간"])),
                    ("근거_면책", b["근거_면책"] if (b["면책기간"] or 0) > 0 else ""),
                    ("적립금 지급", reserve_text(b["적립금지급여부"])),
                    ("근거_적립금", b["근거_적립금"]),
                    ("확인필요", b["확인필요"]),
                    ("별표명", b["별표명"]),
                ]))
    finally:
        reader.quit()
    prods = OrderedDict()
    for r in rows:
        p = prods.setdefault((r["파일명"], r["특약명"]), {"파일명": r["파일명"], "특약명": r["특약명"], "별표명": r["별표명"],
                                                        "급부수": 0, "확인필요": 0})
        p["급부수"] += 1
        p["확인필요"] += 1 if r["확인필요"] else 0
    n_chk = sum(1 for r in rows if r["확인필요"])
    log(f"특약 {len(prods)}개 · 급부 {len(rows)}개 요약 (확인필요 {n_chk}개)")
    return rows, list(prods.values()), status


def run(src, out_dir, log=print):
    """요약 → 엑셀 저장. 엑셀 경로를 돌려줌"""
    rows, prods, status = summarize(src, log)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out = os.path.join(out_dir, f"약관정보_{stamp}.xlsx")
    n = 2
    while os.path.exists(out):
        out = os.path.join(out_dir, f"약관정보_{stamp}_{n}.xlsx")
        n += 1
    cols = ["파일명", "특약명", "급부명", "조항", "지급사유 (언제)", "지급금액 (얼마)", "통계테이블 값", "면책기간", "근거_면책",
            "적립금 지급", "근거_적립금", "확인필요", "별표명"]
    from openpyxl.styles import PatternFill, Alignment, Font
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        pd.DataFrame(rows, columns=cols).to_excel(xw, sheet_name="약관정보", index=False)
        pd.DataFrame(prods, columns=["파일명", "특약명", "별표명", "급부수", "확인필요"]).to_excel(xw, sheet_name="특약목록", index=False)
        pd.DataFrame(status, columns=["파일명", "결과", "급부수"]).to_excel(xw, sheet_name="처리결과", index=False)
        widths = {"지급사유 (언제)": 50, "지급금액 (얼마)": 40, "통계테이블 값": 45, "근거_면책": 40, "확인필요": 35}
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
                w = widths.get(c.value) or min(max(10, max(len(str(x.value or "")) for x in ws[c.column_letter]) * 1.4), 40)
                ws.column_dimensions[c.column_letter].width = w
        ws = xw.book["약관정보"]
        chk = cols.index("확인필요") + 1
        yellow = PatternFill("solid", fgColor=CHECK_FILL)
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if ws.cell(r, chk).value:
                    cell.fill = yellow
    log(f"결과 파일: {out}")
    return out


# =====================================================================
# 화면
# =====================================================================

def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("약관 정보 요약 (급부별 보험금 · 면책기간 · 적립금)")
    root.geometry("780x440")
    root.lift()
    root.attributes("-topmost", True)
    root.after(800, lambda: root.attributes("-topmost", False))
    v_src, v_out = tk.StringVar(), tk.StringVar()
    q = queue.Queue()

    def pick_dir(var):
        d = filedialog.askdirectory()
        if d:
            var.set(d)
            if var is v_src and not v_out.get():
                v_out.set(d)

    def pick_file():
        f = filedialog.askopenfilename(filetypes=[("Word/PDF", "*.docx *.doc *.docm *.pdf")])
        if f:
            v_src.set(f)
            if not v_out.get():
                v_out.set(os.path.dirname(f))

    items = [("약관 폴더/파일", v_src, [("폴더", lambda: pick_dir(v_src)), ("파일", pick_file)]),
             ("결과 저장 위치", v_out, [("폴더", lambda: pick_dir(v_out))])]
    for r, (label, var, btns) in enumerate(items):
        tk.Label(root, text=label, width=15, anchor="w").grid(row=r, column=0, padx=8, pady=6, sticky="w")
        tk.Entry(root, textvariable=var, width=62).grid(row=r, column=1, padx=4, sticky="ew")
        f = tk.Frame(root)
        f.grid(row=r, column=2, padx=6)
        for t, cmd in btns:
            tk.Button(f, text=t, command=cmd, width=5).pack(side="left", padx=2)

    logbox = scrolledtext.ScrolledText(root, height=16)
    logbox.grid(row=3, column=0, columnspan=3, padx=8, pady=8, sticky="nsew")
    root.grid_rowconfigure(3, weight=1)
    root.grid_columnconfigure(1, weight=1)

    def log(m):
        q.put(("log", m))

    def poll():
        try:
            while True:
                kind, m = q.get_nowait()
                if kind == "log":
                    logbox.insert("end", m + "\n")
                    logbox.see("end")
                elif kind == "done":
                    btn.config(state="normal")
                    if messagebox.askyesno("완료", "요약을 만들었습니다. 결과 엑셀을 열까요?"):
                        try:
                            os.startfile(m)
                        except Exception:
                            pass
                elif kind == "error":
                    btn.config(state="normal")
                    messagebox.showerror("오류", m)
        except queue.Empty:
            pass
        root.after(200, poll)

    def start():
        src, out = v_src.get().strip(), v_out.get().strip()
        if not src or not os.path.exists(src):
            messagebox.showwarning("확인", "약관 폴더 또는 파일을 선택해 주세요.")
            return
        if not out or not os.path.isdir(out):
            messagebox.showwarning("확인", "결과 저장 위치를 선택해 주세요.")
            return
        btn.config(state="disabled")
        logbox.delete("1.0", "end")
        log("요약 시작...")

        def work():
            try:
                q.put(("done", run(src, out, log)))
            except Exception as e:
                q.put(("error", str(e)))

        threading.Thread(target=work, daemon=True).start()

    btn = tk.Button(root, text="요약 보기", width=16, command=start)
    btn.grid(row=2, column=0, columnspan=3, pady=6)
    poll()
    root.mainloop()


if __name__ == "__main__":
    run_gui()
